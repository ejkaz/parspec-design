"""Build local, fictional board and investor galleries from an explicit fixture.

No company figures or template assets are packaged. --data must be the v4
sample_data.json supplied with the local design brief. Never use a live model.
"""
from __future__ import annotations

import argparse
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL))
from assets.house import HouseDeck, TEMPLATE_DIR
from assets import shapes as sh
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

D = Decimal
BOARD_IDS = ['BD01','BD01a']+[f'BD{i:02}' for i in range(2,17)]
INVESTOR_IDS = [f'IN{i:02}' for i in range(1,20)]
PICKS = [('BD01','bb'),('BD01a','ba'),('BD02','ba'),('BD03','ba'),
         ('BD04','bb'),('BD05','ba'),('BD06','ba'),('BD07','bb'),
         ('BD08','bb'),('BD09','ba'),('BD10','bb'),('BD11','ba'),
         ('BD12','bb'),('BD13','ba'),('BD14','ba'),('BD15','bb')] + [
         (f'IN{i:02}', 'ia' if i in (2,6,13) else 'ib') for i in range(1,15)]


def q(value, places=2):
    return D(str(value)).quantize(D(10) ** -places, rounding=ROUND_HALF_UP)


def num(value, places=1):
    v = q(value,places)
    return f'({abs(v):.{places}f})' if v < 0 else f'{v:.{places}f}'


def money(value, places=1):
    return f'${num(value,places)}M' if value >= 0 else f'(${num(-value,places)}M)'


def percent(value,places=0):
    return num(D(str(value))*100,places)+'%'


def ratio(value,places=1):
    return num(value,places)+'×'


def month_label(raw):
    month, yy = raw.split()
    year = 2000 + int(yy)
    m = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'].index(month)+1
    fy = year if m >= 2 else year-1
    quarter = ((m-2)%12)//3+1
    return f'{month} {year}\n(Q{quarter} FY{str(fy)[2:]})'


def calculations(d):
    b=d['arr_bridge_q3_fy26']; p=d['pnl']; fy=d['fy27_plan']
    # Store the quarter at two decimals, sum those stored quarters, then display.
    gp=[q(r*g) for r,g in zip(p['revenue'],p['gross_margin'])]
    ebitda=[q(g-o) for g,o in zip(gp,p['opex'])]
    gp27=[q(r*g) for r,g in zip(fy['revenue'],fy['gross_margin'])]
    weighted=[q(s['amount']*s['win_rate']) for s in d['pipeline_q4_fy26']['stages']]
    hist=d['arr_history']['arr']; peers=sorted(c['valuation_arr'] for c in d['comps'])
    median=(peers[3]+peers[4])/2
    def percentile(frac):
        rank=D(len(peers)-1)*frac; low=int(rank)
        return peers[low]+(rank-low)*(peers[min(low+1,len(peers)-1)]-peers[low])
    return dict(gp=gp,ebitda=ebitda,gp27=gp27,
                cor=[r-g for r,g in zip(p['revenue'],gp)],
                net=b['end']-b['start'], gains=b['new']+b['expansion'],
                losses=b['contraction']+b['churn'],
                weighted=weighted,weighted_total=sum(weighted),
                coverage=sum(weighted)/d['pipeline_q4_fy26']['target_net_new'],
                median=median,p25=percentile(D('.25')),p75=percentile(D('.75')),
                median_growth=sum(sorted(c['growth'] for c in d['comps'])[3:5])/2,
                multiple=d['round']['pre_money']/b['end'],
                growth=b['end']/hist[4]-1,two_year=b['end']/hist[0],
                arr=hist,net_new=[hist[i]-hist[i-1] for i in range(6,len(hist))])


class Gallery:
    def __init__(self,d,sha):
        self.d,self.c,self.sha=d,calculations(d),sha
        self.deck=HouseDeck('investor')
        self.i,self.ch=self.deck.investor,self.deck.charts
        self.ids=[]

    def slide(self,sid,eyebrow='',title='',subtitle=''):
        s=self.deck.slide('plain',keep=()); self.ids.append(sid)
        self.deck.notes(s, f'Illustrative sample data (template fixture v4).\n'
                        f'Parspec, Inc. standalone.\nSlide: {sid}\nFixture SHA-256: {self.sha}')
        y=self.i.header(s,eyebrow,title,subtitle) if title else 184
        return s,y

    def table(self,s,headers,rows,*,x=56,y=220,w=1168,h=350,**kw):
        kw.setdefault('total_rows',[j for j,row in enumerate(rows) if row[0] in ('Total','EBITDA','Gross profit')])
        return self.ch.fin_table(s,headers,rows,x=x,y=y,w=w,h=h,**kw)

    def cell_lines(self,cell,lines,*,padding=18,left=None,fill='card'):
        """Native editable paragraphs in an unthemed table cell."""
        cell.fill.solid();cell.fill.fore_color.rgb=self.i.color(fill)
        cell.margin_left=Inches((padding if left is None else left)/128)
        cell.margin_right=Inches(padding/128)
        cell.margin_top=cell.margin_bottom=Inches(padding/128)
        cell.vertical_anchor=MSO_ANCHOR.TOP
        props=cell._tc.get_or_add_tcPr()
        for side in ('lnL','lnR','lnT','lnB'):
            edge=OxmlElement('a:'+side);edge.set('w','6350')
            solid=OxmlElement('a:solidFill');clr=OxmlElement('a:srgbClr');clr.set('val',str(self.i.color('line')))
            solid.append(clr);edge.append(solid);props.append(edge)
        cell.text=''
        for j,(text,size,color,bold) in enumerate(lines):
            para=cell.text_frame.paragraphs[0] if j==0 else cell.text_frame.add_paragraph()
            para.space_before=Pt(0);para.space_after=self.ch.pt(8 if j==0 else 12 if j<len(lines)-1 else 0)
            para.line_spacing=self.ch.pt(size*1.15);para.alignment=PP_ALIGN.LEFT
            run=para.add_run();run.text=str(text);run.font.name=self.deck.brand.font
            run.font.size=self.ch.pt(size);run.font.bold=bold;run.font.color.rgb=self.i.color(color)

    @staticmethod
    def plot_layout(chart,x,y,w,h):
        area=chart._chartSpace.chart.plotArea
        layout=area.find(qn('c:layout'))
        if layout is None:
            layout=OxmlElement('c:layout');area.insert(0,layout)
        for child in list(layout):layout.remove(child)
        manual=OxmlElement('c:manualLayout')
        for name,val in [('layoutTarget','inner'),('xMode','edge'),('yMode','edge'),('wMode','factor'),('hMode','factor'),('x',x),('y',y),('w',w),('h',h)]:
            elem=OxmlElement('c:'+name);elem.set('val',str(val));manual.append(elem)
        layout.append(manual)

    def customer_outcomes(self):
        d=self.d
        s,y=self.slide('IN15','Customer outcomes','Customers Quote Faster and Win More',
                       'Illustrative customer outcomes with fiscal periods and attribution placeholders.')
        rows=[[o['customer']+'\n'+o['segment'],o['metric']+'\n'+o['before']+' → '+o['after'],o['period'],o['status'],o['attribution']] for o in d['customer_outcomes']]
        native=self.table(s,['Customer','Outcome before → after','Period','Status','Attribution'],rows,y=222,h=390,
                          widths=[.25,.27,.19,.10,.19],font_px=18,numeric_cols=[])
        native.name='native:customer_outcomes'
        # A single rule separates each evidence row. No copied logos or quotes.
        for row in list(native.table.rows)[1:]:
            for cell in row.cells:
                props=cell._tc.get_or_add_tcPr();edge=OxmlElement('a:lnB');edge.set('w','6350')
                fill=OxmlElement('a:solidFill');clr=OxmlElement('a:srgbClr');clr.set('val',str(self.i.color('line')))
                fill.append(clr);edge.append(fill);props.append(edge)

    def customer_value(self):
        value=self.d['customer_value'];total=sum(b['value_k'] for b in value['benefits'])
        s,y=self.slide('IN16','Customer value',f'${num(total,0)}K of Annual Value for a 12-Branch Distributor',
                       value['customer']+'. '+value['horizon']+' scope.')
        rows=[[b['benefit'],b['mechanism'],num(b['value_k']),b['status']] for b in value['benefits']]
        rows.append(['Total annual value','',num(total),'Mixed'])
        native=self.table(s,['Benefit','Mechanism','Annual value · US$K','Status'],rows,y=226,h=330,
                          widths=[.24,.43,.18,.15],font_px=20,numeric_cols=[2],total_rows=[3])
        native.name='native:customer_value'
        for j,benefit in enumerate(value['benefits'],1):
            cell=native.table.cell(j,3);cell.fill.fore_color.rgb=self.i.color('band' if benefit['status']=='Measured' else 'forecast_fill')
            for para in cell.text_frame.paragraphs:
                for run in para.runs:run.font.color.rgb=self.i.color('text' if benefit['status']=='Measured' else 'accent')
        self.i.text(s,56,584,1168,64,ratio(total/value['parspec_cost_k'])+' the annual subscription',size=44,bold=True)
        self.i.text(s,56,646,1168,26,'Annual subscription: $'+num(value['parspec_cost_k'])+'K',size=16,color='muted')

    def team(self):
        s,y=self.slide('IN17','Team','[Leadership team]','Names, credentials and portraits remain template placeholders.')
        native=s.shapes.add_table(2,3,Inches(56/128),Inches(224/128),Inches(1168/128),Inches(416/128))
        native.name='native:leadership_grid';native.table.first_row=False;native.table.horz_banding=False
        cw,ch=1168/3,208
        for j,person in enumerate(self.d['team']):
            xx=56+(j%3)*cw;yy=224+(j//3)*ch
            self.cell_lines(native.table.cell(j//3,j%3),[(person['name'],24,'text',True),
                            (person['role'],18,'text',False),(person['credential'],14,'muted',False)],padding=22,left=132,fill='ground')
            circle=sh.shape(s,MSO_SHAPE.OVAL,*[v/128 for v in (xx+24,yy+24,80,80)],fill=self.i.color('band'),line=self.i.color('line'),lw=.56)
            circle.name=f'portrait-placeholder:{j+1}'
            self.i.text(s,xx+24,yy+54,80,24,'Portrait',size=14,color='muted',align='c')

    def product_mechanism(self):
        mechanism=self.d['product_mechanism']
        s,y=self.slide('IN18','Product mechanism','From Specification to Purchase Order',
                       'Specifications become quotes, submittals and purchase orders.')
        for k,group in enumerate(('inputs','engine','outputs')):
            yy=216+k*154
            native=s.shapes.add_table(1,5,Inches(56/128),Inches(yy/128),Inches(1168/128),Inches(104/128))
            native.name='native:mechanism_'+group;native.table.first_row=False;native.table.horz_banding=False
            native.table.columns[0].width=Inches(140/128)
            for c in range(1,5):native.table.columns[c].width=Inches((1028/4)/128)
            self.cell_lines(native.table.cell(0,0),[(group.upper(),18,'accent',True)],padding=18)
            for j,item in enumerate(mechanism[group],1):
                status=mechanism['status'].get(item,'Live')
                self.cell_lines(native.table.cell(0,j),[(item,18,'text',True),(status,14,'accent' if status=='Planned' else 'muted',False)],padding=18,
                                fill='forecast_fill' if status=='Planned' else 'card')
            if k<2:
                connector=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(640/128),Inches((yy+104)/128),Inches(640/128),Inches((yy+154)/128))
                connector.name=f'mechanism-connector:{group}'
                connector.line.color.rgb=self.i.color('line');connector.line.width=Pt(1.12)
                line=connector._element.spPr.get_or_add_ln();arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');line.append(arrow)

    def cash_bridge(self,sid):
        bridge=self.d['ebitda_to_cash_q3_fy26']
        keys=['ebitda','working_capital','deferred_revenue','capex','interest','other','net_burn']
        labels=['EBITDA','Working capital','Deferred revenue','Software & capex','Interest','Other','Net burn']
        s,y=self.slide(sid,'EBITDA to cash',f'Q3 FY26 Net Burn {money(-bridge["net_burn"])} vs EBITDA $({num(-bridge["ebitda"],2)})M',
                       'Operating result to net cash movement. Parspec, Inc. standalone.')
        self.ch.waterfall(s,labels,[bridge[k] for k in keys],x=56,y=232,w=1168,h=348,
                          minimum=-3.7,maximum=0,label_above=True,label_decimals=2,number_format='0.00;(0.00);0.00')
        self.i.text(s,56,194,1168,24,'US$M · Q3 FY26',kind='label',color='muted')
        self.i.callout(s,'Net burn '+money(-bridge['net_burn'])+' ties to the board plan-versus-actual scorecard.',y=614,h=40)

    def strip(self,s,items,*,y=564,x=56,w=1168):
        cw=w/len(items)
        self.i.rule(s,x,y,w)
        for k,(label,value,context) in enumerate(items):
            xx=x+k*cw
            if k:self.i.rect(s,xx,y,1,74,fill='line')
            self.i.text(s,xx+16,y+12,cw-32,22,label.upper(),kind='label',color='muted')
            self.i.text(s,xx+16,y+38,cw-32,34,value,size=28,bold=True)
            if context:self.i.text(s,xx+16,y+74,cw-32,24,context,size=14,color='muted')

    def decisions(self,s,*,x,y,w,compact=False):
        ds=self.d['decisions']; rh=120 if compact else 124
        self.i.rect(s,x,y,w,rh*len(ds)+48,fill='card' if compact else 'ground')
        self.i.rule(s,x,y,w,h=4,color='accent')
        for j,decision in enumerate(ds):
            yy=y+22+j*rh
            self.i.text(s,x+18,yy,38,38,str(j+1),size=30,color='accent',bold=True)
            self.i.text(s,x+68,yy,w-88,82,decision['ask'],size=16 if compact else 24,bold=True)
            self.i.text(s,x+68,yy+86,w-88,24,
                        decision['owner']+' · '+decision['by'],size=14,color='muted')
            if j<len(ds)-1:self.i.rule(s,x+68,yy+114,w-88)

    def board(self):
        d,c=self.d,self.c; b=d['arr_bridge_q3_fy26']; p=d['pnl']; cash=d['cash_runway']; hc=d['headcount']
        s,_=self.slide('BD01'); self.i.cover(s,board=True,partner='BOARD OF DIRECTORS'); self.arcs(s,small=True)

        s,y=self.slide('BD01a','Agenda',"Today's Agenda and Board Asks",f"{sum(v['minutes'] for v in d['agenda'])} minutes. All three decisions at this meeting.")
        rows=[[str(j+1),a['item'],a['owner'],str(a['minutes'])] for j,a in enumerate(d['agenda'])]
        rows.append(['','Total','',str(sum(a['minutes'] for a in d['agenda']))])
        self.table(s,['#','Item','Owner','Minutes'],rows,x=56,y=232,w=710,h=348,widths=[.06,.62,.17,.15],font_px=18,numeric_cols=[3])
        self.decisions(s,x=808,y=232,w=416,compact=True)

        s,y=self.slide('BD02','Quarter in review',f'ARR Up {money(c["net"])} to {money(b["end"])}; Burn Under Plan','Q3 FY26 against the board plan. Parspec, Inc. standalone.')
        vals=[money(b['end']),money(c['net']),percent(d['unit_economics']['ndr']),percent(p['gross_margin'][2]),money(d['plan_vs_actual_q3_fy26'][4]['actual']),money(d['plan_vs_actual_q3_fy26'][5]['actual'])]
        labels=['Ending ARR','Net new ARR','NDR · trailing year','Gross margin','Net burn','Cash']
        contexts=[f'+{money(c["net"])} QoQ',f'{money(c["net"])} vs {money(d["plan_vs_actual_q3_fy26"][0]["plan"]-b["start"])} plan',d['quarter_kpis'][2]['delta'],d['quarter_kpis'][3]['delta'],d['quarter_kpis'][4]['delta'],d['quarter_kpis'][5]['delta']]
        actuals=d['plan_vs_actual_q3_fy26']
        contexts=[f'+{money(c["net"])} ▲ QoQ',
                  f'−{money(actuals[0]["plan"]-b["end"])} ▼ vs plan',
                  d['quarter_kpis'][2]['delta']+' ▲',
                  d['quarter_kpis'][3]['delta']+' ▲',
                  f'−{money(actuals[4]["plan"]-actuals[4]["actual"])} ▲ vs plan',
                  f'+{money(actuals[5]["actual"]-actuals[5]["plan"])} ▲ vs plan']
        grid=s.shapes.add_table(2,3,Inches(56/128),Inches(224/128),Inches(752/128),Inches(372/128))
        grid.name='native:kpi_grid'; grid.table.first_row=False;grid.table.horz_banding=False
        for j in range(6):
            cell=grid.table.cell(j//3,j%3)
            self.cell_lines(cell,[(labels[j].upper(),14,'muted',True),
                                  (vals[j],42,'text',True),
                                  (contexts[j],16,'muted',False)],padding=22)
        self.i.rule(s,56,224,752,color='accent',h=3)
        self.i.text(s,856,224,368,28,'WHAT CHANGED',kind='eyebrow',color='accent')
        changes=[('ARR gains',f'New logos {money(b["new"])} plus expansion {money(b["expansion"])}.'),
                 ('ARR losses',f'Contraction {money(-b["contraction"])} and churn {money(-b["churn"])}. Net new {money(c["net"])} vs {money(actuals[0]["plan"]-b["start"])} plan.'),
                 ('Cash and spend',f'Cash {money(actuals[5]["actual"])} is {money(actuals[5]["actual"]-actuals[5]["plan"])} ahead. OpEx {money(p["opex"][2])} vs {money(actuals[3]["plan"])} plan.')]
        for j,(label,body) in enumerate(changes):
            yy=272+j*118
            self.i.text(s,856,yy,368,25,label,size=20,bold=True)
            self.i.text(s,856,yy+34,368,82,body,size=18)
            if j<2:self.i.rule(s,856,yy+108,368)

        s,y=self.slide('BD03','ARR bridge',f'{money(c["net"])} Net New ARR in Q3 FY26','Starting ARR to ending ARR. Parspec, Inc. standalone.')
        self.ch.waterfall(s,['Start\nQ3 FY26','New logos','Expansion','Contraction','Churn','End\nQ3 FY26'],[b[k] for k in ('start','new','expansion','contraction','churn','end')],x=56,y=214,w=1168,h=330,minimum=15,maximum=23,axis_start_label='Axis starts at $15M',label_above=True)
        self.strip(s,[('New + expansion',money(c['gains']),''),('Contraction + churn',money(c['losses']),''),('Net new ARR',money(c['net']),''),('GRR · trailing year',percent(d['unit_economics']['grr']),''),('NDR · trailing year',percent(d['unit_economics']['ndr']),'')],y=566)

        s,y=self.slide('BD04','Growth',f'Plan Reaches ${num(c["arr"][-1],0)}M Ending ARR in FY27','Quarterly ending ARR. Parspec, Inc. standalone.')
        self.ch.act_forecast_chips(s,'ACTUAL · THROUGH Q3 FY26','FORECAST · FROM Q4 FY26',x=230,y=184,w=720)
        self.ch.column_actual_forecast(s,d['arr_history']['quarters'][6:],c['arr'][6:],3,x=56,y=222,w=1168,h=330)
        self.strip(s,[('Q3 FY26 ARR',money(b['end']),f'+{percent(c["growth"])} year on year'),('FY27 exit ARR · plan',money(c['arr'][-1]),f'+{percent(c["arr"][-1]/c["arr"][9]-1)} on FY26 exit'),('FY27 net new ARR · plan',money(c['arr'][-1]-c['arr'][9]),'Over four fiscal quarters')],y=570)

        s,y=self.slide('BD05','Board plan vs actual','Ahead on Cash and Burn, Behind on ARR','Q3 FY26 Board View. Parspec, Inc. standalone.')
        rows=[]
        for j,r in enumerate(d['plan_vs_actual_q3_fy26']):
            delta=r['actual']-r['plan']; good=delta>0 if j in (0,1,2,5) else delta<0
            vs=('+' if delta>0 else '−')+(num(abs(delta)*100,0)+' pt' if j==2 else num(abs(delta)))
            rows.append([r['metric'],percent(r['plan']) if j==2 else num(r['plan']),percent(r['actual']) if j==2 else num(r['actual']),vs,('▲ Ahead' if good else '▼ Behind')])
        self.table(s,['Metric · Q3 FY26','Board plan','Actual','Variance','Status'],rows,y=210,h=350,widths=[.36,.16,.16,.16,.16],font_px=20)
        self.i.callout(s,f'Cash finished {money(d["plan_vs_actual_q3_fy26"][5]["actual"]-d["plan_vs_actual_q3_fy26"][5]["plan"])} ahead of plan; ending ARR finished {money(d["plan_vs_actual_q3_fy26"][0]["plan"]-b["end"])} behind on net new.',y=590)

        s,y=self.slide('BD06','P&L','Gross Margin Plan Rises Across FY26','Q4 FY26 forecast. FY26 total: actual + forecast. Parspec, Inc. standalone.')
        rows=self.pnl_rows(detail=False)
        self.table(s,['US$M']+p['quarters']+['FY26'],rows,y=230,h=320,widths=[.29]+[.142]*5,statuses=['','Actual','Actual','Actual','Forecast','Mixed'],highlight_col=4,font_px=20)
        self.i.callout(s,f'Revenue rises {money(p["revenue"][-1]-p["revenue"][0])} across FY26; EBITDA loss narrows from {money(c["ebitda"][0])} to {money(c["ebitda"][-1])}.',y=588)

        s,y=self.slide('BD07','Cash',f'{cash["runway_months_without_debt"]} Months of Runway Before the Debt Draw','Month-end cash forecast. Parspec, Inc. standalone.')
        labels=['Oct']+[m.split()[0] for m in cash['months']]
        values=[d['plan_vs_actual_q3_fy26'][5]['actual']]+cash['cash_end']
        self.ch.act_forecast_chips(s,'ACTUAL · OCT 2026 (Q3 FY26)','FORECAST · FROM NOV 2026 (Q4 FY26)',x=84,y=186,w=720)
        chart=self.ch.column_actual_forecast(s,labels,values,1,x=56,y=280,w=820,h=252,font_px=14,legend=False)
        self.plot_layout(chart.chart,.075,.06,.90,.80)
        # Month labels stay single-line; groups identify the fiscal year/quarter.
        plot_x,plot_w=56+820*.075,820*.90
        for first,count,fiscal in [(0,1,'Q3 FY26'),(1,3,'Q4 FY26'),(4,3,'Q1 FY27'),(7,3,'Q2 FY27'),(10,3,'Q3 FY27')]:
            xx=plot_x+plot_w*first/len(labels);ww=plot_w*count/len(labels)
            self.i.rect(s,xx,250,ww,26,fill='band')
            band=self.i.text(s,xx-12,252,ww+24,22,fiscal,size=14,color='muted',align='c')
            band.name=f'fiscal-band:{first}:{count}:{fiscal}'
        self.i.text(s,56,540,820,22,'Oct 2026 to Oct 2027 · fiscal quarters shown above',size=14,color='muted')
        self.i.text(s,362,218,500,24,f'{money(cash["debt_draw"]["amount"])} debt draw · Feb 2027 (Q1 FY27)',kind='label',color='accent')
        self.table(s,['US$M','Next twelve months','Monthly average'],[['Net burn',num(sum(cash['net_burn'])),num(sum(cash['net_burn'])/len(cash['net_burn']))]],x=56,y=562,w=820,h=66,font_px=16,widths=[.3,.35,.35])
        self.i.tile(s,912,184,312,132,'Cash · Q3 FY26',money(d['plan_vs_actual_q3_fy26'][5]['actual']),size=42)
        self.i.tile(s,912,336,312,132,'Net burn · Q3 FY26',money(d['plan_vs_actual_q3_fy26'][4]['actual']),size=42)
        self.i.tile(s,912,488,312,166,'Runway without draw',f'{cash["runway_months_without_debt"]} months',f'{cash["runway_months_with_debt"]} months with draw',size=38)

        s,y=self.slide('BD08','Scenarios','Cash Positive Through FY29 in Board and Bank\nCases','Named case outcomes. Parspec, Inc. standalone.')
        self.ch.bar_ranked(s,[a['case'] for a in d['cases']],[a['min_cash'] for a in d['cases']],x=56,y=y+35,w=440,h=280)
        rows=[[a['case'],num(a['fy27_exit_arr']),num(a['fy27_revenue']),a['cash_out']] for a in d['cases']]
        self.table(s,['Case','FY27 exit ARR','FY27 revenue','Cash-out'],rows,x=520,y=y+20,w=704,h=280,font_px=18,widths=[.15,.22,.22,.41])
        self.i.text(s,56,y,440,28,'MINIMUM CASH · US$M',kind='label',color='muted')
        self.i.callout(s,f'The Bank case holds {money(d["cases"][1]["min_cash"])} at its low point; only the Bear case runs out of cash, in {d["cases"][2]["cash_out"]}.',y=592)

        s,y=self.slide('BD09','Headcount',f'Net +{sum(hc["end"])-sum(hc["start"])} Hires to {sum(hc["end"])}; Q4 FY26 Plan Is {sum(hc["plan_q4"])}','One hiring plan. Parspec, Inc. standalone.')
        self.ch.waterfall(s,['Start\nQ3 FY26','Joiners','Leavers','End\nQ3 FY26','Plan\nQ4 FY26'],[sum(hc['start']),sum(hc['joiners']),-sum(hc['leavers']),sum(hc['end']),sum(hc['plan_q4'])],total_indices=(0,3,4),plan_index=4,number_format='0',x=56,y=244,w=490,h=306,minimum=90,maximum=112,axis_start_label='Axis starts at 90 heads',label_above=True)
        rows=[[f,str(hc['start'][j]),'+'+str(hc['joiners'][j]),'('+str(hc['leavers'][j])+')' if hc['leavers'][j] else '—',str(hc['end'][j]),str(hc['plan_q4'][j])] for j,f in enumerate(hc['functions'])]
        rows.append(['Total',str(sum(hc['start'])),'+'+str(sum(hc['joiners'])),'('+str(sum(hc['leavers']))+')',str(sum(hc['end'])),str(sum(hc['plan_q4']))])
        self.table(s,['Function','Start','Joiners','Leavers','End','Q4 FY26\nPlan'],rows,x=570,y=220,w=654,h=330,widths=[.36,.12,.13,.13,.12,.14],font_px=16,highlight_col=5)
        self.i.callout(s,f'The Q4 FY26 plan adds {sum(hc["plan_q4"])-sum(hc["end"])} heads to reach {sum(hc["plan_q4"])}.',y=590)

        pipe=d['pipeline_q4_fy26']; stages=pipe['stages']
        s,y=self.slide('BD10','Pipeline',f'{ratio(c["coverage"])} Coverage on the Q4 FY26 Net New Target','Weighted pipeline and the next-quarter target. Parspec, Inc. standalone.')
        self.ch.stacked_bar(s,['Weighted pipeline'],[dict(name=t['stage'],values=[c['weighted'][j]],color=f'chart_{j+1}') for j,t in enumerate(stages)],x=56,y=200,w=900,h=126,maximum=5)
        self.i.text(s,1008,220,216,64,money(c['weighted_total']),kind='stat')
        target_x=228+float(pipe['target_net_new']/5)*692
        self.i.rect(s,target_x,200,2,66,fill='text')
        self.i.text(s,target_x+10,184,400,24,f'Q4 FY26 target · {money(pipe["target_net_new"])}',kind='label')
        self.i.text(s,56,334,1168,26,f'Q4 FY26 net new target {money(pipe["target_net_new"])}',size=18,color='muted')
        rows=[[t['stage'],num(t['amount']),percent(t['win_rate']),num(c['weighted'][j])] for j,t in enumerate(stages)]
        rows.append(['Total',num(sum(t['amount'] for t in stages)),'—',num(c['weighted_total'])])
        self.table(s,['Stage · US$M','Open pipeline','Win rate','Weighted'],rows,y=376,h=222,widths=[.4,.2,.2,.2],font_px=18)
        self.i.text(s,56,622,1168,28,f'Proposal + Commit: {money(sum(c["weighted"][2:]))}',size=18)

        s,y=self.slide('BD11','For the board','Three Decisions Today','All decisions are for this meeting. Parspec, Inc. standalone.')
        self.decisions(s,x=56,y=202,w=1168)

        s,y=self.slide('BD12','Appendix','P&L Detail by Quarter','Q4 FY26 forecast. FY26 total: actual + forecast. Parspec, Inc. standalone.')
        self.table(s,['US$M']+p['quarters']+['FY26'],self.pnl_rows(detail=True),y=226,h=316,widths=[.29]+[.142]*5,statuses=['','Actual','Actual','Actual','Forecast','Mixed'],highlight_col=4,font_px=16)
        self.i.text(s,56,568,560,22,'DEFINITIONS AND ROUNDING',kind='label',color='accent')
        self.i.text(s,56,596,560,62,'Store two decimals half-up. Display $M at one decimal half-up. Totals sum stored quarters; margins use totals. Runway and case cash-out are supplied model assertions.',size=13,color='muted')
        self.i.text(s,660,568,564,22,'SOURCES',kind='label',color='accent')
        self.i.text(s,660,596,564,62,'Illustrative sample data (template fixture v4). Company NDR / GRR are trailing-twelve-month assertions. ARR bridge is quarterly. Minimum-cash horizon is unspecified.',size=13,color='muted')

        s,y=self.slide('BD13','For approval',f'Authorize the ${num(cash["debt_draw"]["amount"],0)}M Debt Draw Window','Draw window: Feb 2027 (Q1 FY27). Parspec, Inc. standalone.')
        self.i.tile(s,56,184,366,136,'Draw amount',money(cash['debt_draw']['amount']),forecast=True,size=42)
        self.i.tile(s,446,184,366,136,'Draw month','Feb 2027','Q1 FY27',forecast=True,size=38)
        self.i.tile(s,836,184,388,136,'Runway · months',f'{cash["runway_months_without_debt"]} to {cash["runway_months_with_debt"]}',size=42)
        without=[v-(cash['debt_draw']['amount'] if j>=3 else 0) for j,v in enumerate(cash['cash_end'])]
        rows=[['Cash · Jan 2027 (Q4 FY26)',num(without[2]),num(cash['cash_end'][2])],['Cash · Feb 2027 (Q1 FY27)',num(without[3]),num(cash['cash_end'][3])],['Cash · Oct 2027 (Q3 FY27)',num(without[-1]),num(cash['cash_end'][-1])],['Net burn · next twelve months',num(sum(cash['net_burn'])),num(sum(cash['net_burn']))]]
        self.table(s,['US$M','Without draw','With draw'],rows,x=56,y=358,w=730,h=222,widths=[.56,.22,.22],font_px=17,highlight_col=2)
        self.table(s,['Months','Without draw','With draw'],[['Runway',str(cash['runway_months_without_debt']),str(cash['runway_months_with_debt'])]],x=56,y=598,w=730,h=54,widths=[.56,.22,.22],font_px=16,highlight_col=2)
        self.i.callout(s,d['decisions'][1]['ask']+'\nCFO · This meeting',x=822,y=358,w=402,h=254,ask=True)

        s,y=self.slide('BD14','Risks','Three Risks We Are Managing','Named owners and the mitigation in flight. Parspec, Inc. standalone.')
        rows=[]
        for j,r in enumerate(d['risks']):
            mitigation=r['mitigation'].replace(' Q1',' Q1 FY27').replace(' Q2',' Q2 FY27')
            rows.append([str(j+1),r['risk'],mitigation,r['owner']])
        self.table(s,['#','Risk','Mitigation in flight','Owner'],rows,y=212,h=378,widths=[.06,.4,.46,.08],font_px=18,numeric_cols=[])

        s,y=self.slide('BD15'); self.deck.logo(s,x=56/128,y=48/128,w=176/128)
        self.i.text(s,270,52,850,30,'Q3 FY26 BOARD MEETING',kind='label',color='muted')
        self.i.text(s,56,278,1168,28,'AGENDA ITEM 2 · 20 MINUTES',kind='eyebrow',color='accent')
        self.i.text(s,56,322,1168,90,'Financials',kind='cover')
        self.i.rule(s,56,440,192,color='accent',h=3)
        self.i.text(s,56,480,760,26,'IN THIS SECTION',kind='label',color='accent')
        self.i.text(s,56,518,850,36,'P&L · Cash & runway · Scenarios · Headcount',size=22,color='muted')
        self.i.text(s,940,480,284,26,'PRESENTED BY',kind='label',color='accent')
        self.i.text(s,940,518,284,36,'CFO · 20 minutes',size=22)
        self.cash_bridge('BD16')
        return self.deck

    def pnl_rows(self,detail=False):
        p,c=self.d['pnl'],self.c
        values=[('Revenue',p['revenue']),('Gross profit',c['gp']),('Gross margin',[g/r for g,r in zip(c['gp'],p['revenue'])]),('Operating expenses',p['opex']),('EBITDA',c['ebitda'])]
        if detail:
            values=[('Revenue',p['revenue']),('Revenue growth, QoQ',[None]+[p['revenue'][j]/p['revenue'][j-1]-1 for j in range(1,4)]),('Cost of revenue',c['cor']),('Gross profit',c['gp']),('Gross margin',[g/r for g,r in zip(c['gp'],p['revenue'])]),('Operating expenses',p['opex']),('EBITDA',c['ebitda']),('EBITDA margin',[e/r for e,r in zip(c['ebitda'],p['revenue'])])]
        rows=[]
        for label,vs in values:
            is_pct='margin' in label or 'growth' in label
            if 'growth' in label:total=None
            elif label=='Gross margin':total=sum(c['gp'])/sum(p['revenue'])
            elif label=='EBITDA margin':total=sum(c['ebitda'])/sum(p['revenue'])
            else:total=sum(vs)
            rows.append([label]+[('—' if v is None else percent(v,1 if detail and 'growth' in label else 0) if is_pct else num(v,2 if detail else 1)) for v in list(vs)+[total]])
        return rows

    def arcs(self,s,small=False):
        # Editable circles reproduce the picked cover arcs, clipped by canvas.
        cx,cy=1312,780
        for dd,fill,line in [(472,'forecast_fill',None),(552,None,'accent'),(616,None,'accent')]:
            if small:dd*=.72
            sh.shape(s,MSO_SHAPE.OVAL,*[v/128 for v in (cx-dd/2,cy-dd/2,dd,dd)],fill=self.i.color(fill) if fill else None,line=self.i.color(line) if line else None,lw=.56)

    def investor(self):
        d,c=self.d,self.c; b=d['arr_bridge_q3_fy26']; u=d['unit_economics']; r=d['round']; market=d['market']
        s,_=self.slide('IN01'); self.i.cover(s); self.arcs(s)
        s,y=self.slide('IN02','Why Parspec',f'{money(b["end"])} ARR, {percent(u["ndr"]) } NDR, {percent(d["pnl"]["gross_margin"][2])} Gross Margin','Q3 FY26. Parspec, Inc. standalone.')
        tiles=[('Scale',money(b['end']),f'{ratio(c["two_year"])} the ARR in Q3 FY24'),('Durability',percent(u['ndr']),f'{percent(u["grr"])} GRR · trailing twelve months'),('Efficiency',percent(d['pnl']['gross_margin'][2]),f'Up {num((u["gross_margin_trend"]["gm"][-1]-u["gross_margin_trend"]["gm"][0])*100,0)} pts from Q1 FY25'),('FY27 plan',money(c['arr'][-1]),'Ending ARR · Q4 FY27')]
        for j,(label,value,ctx) in enumerate(tiles):self.i.tile(s,56+j*298,224,274,276,label,value,ctx,forecast=j==3,size=52)
        self.i.callout(s,f'${num(market["tam"],0)}B US market; ${num(market["sam"])}B serviceable today; ${num(market["som"]*1000,0)}M five-year opportunity.',y=552,h=88)

        s,y=self.slide('IN03','Growth',f'ARR Up {ratio(c["two_year"])} in Two Years','Quarterly ending ARR. Parspec, Inc. standalone.')
        self.ch.act_forecast_chips(s,'ACTUAL · THROUGH Q3 FY26','FORECAST · FROM Q4 FY26',x=480,y=190,w=740)
        self.ch.column_actual_forecast(s,d['arr_history']['quarters'],c['arr'],9,x=56,y=240,w=1168,h=350)
        self.i.rule(s,124,620,638,color='muted')
        self.i.rect(s,124,612,1,9,fill='muted'); self.i.rect(s,762,612,1,9,fill='muted')
        self.i.text(s,124,632,850,26,f'Q3 FY24 to Q3 FY26 · {ratio(c["two_year"])} actual ARR',size=18,color='muted')

        s,y=self.slide('IN04','Retention',f'Mature Cohorts Expand to {num(min(row[-1] for row in [[v for v in r if v is not None] for r in d["cohort_ndr"]["ndr"]] if len(row)>1)*100,0)}–{num(max(v for row in d["cohort_ndr"]["ndr"] for v in row if v is not None)*100,0)}% of Year-1 ARR','Cumulative cohort expansion. Parspec, Inc. standalone.')
        self.i.text(s,56,254,330,118,percent(u['ndr']),kind='hero')
        self.i.text(s,56,378,330,64,'Net dollar retention\nTrailing twelve months',size=20,bold=True)
        self.i.rule(s,56,470,280)
        self.i.text(s,56,498,280,68,percent(u['grr']),kind='stat')
        self.i.text(s,56,576,280,32,'Gross retention',size=20,color='muted')
        co=d['cohort_ndr']; self.ch.heatmap_table(s,co['cohorts'],co['years'],co['ndr'],x=450,y=248,w=774,h=334)

        s,y=self.slide('IN05','Unit economics',f'{u["cac_payback_months"]}-Month CAC Payback at {percent(u["gross_margin_trend"]["gm"][-1])} Gross Margin','Q3 FY26 efficiency. Parspec, Inc. standalone.')
        self.i.text(s,56,240,330,130,str(u['cac_payback_months']),kind='hero')
        self.i.text(s,56,378,330,54,'Months CAC payback',size=22,bold=True)
        self.i.rule(s,56,454,280)
        self.i.text(s,56,480,310,44,ratio(u['ltv_cac'])+' LTV / CAC',size=28,bold=True)
        self.i.text(s,56,542,310,60,percent(u['ndr'])+' NDR\n'+percent(u['grr'])+' GRR',size=20,color='muted')
        gm=u['gross_margin_trend']; self.ch.line_trend(s,gm['quarters'],[dict(name='Gross margin',values=gm['gm'],color='text')],x=412,y=250,w=812,h=340,number_format='0%',minimum=.6,maximum=.8)
        self.i.text(s,440,616,740,28,f'+{num((gm["gm"][-1]-gm["gm"][0])*100,0)} pts since Q1 FY25',size=20,color='muted')

        s,y=self.slide('IN06','Market',f'${num(market["tam"],0)}B Market, ${num(market["sam"])}B Serviceable Today','US electrical distribution software spend. Parspec, Inc. metrics are standalone.')
        # Source-picked nested areas, proportional to TAM, are editable circles.
        import math
        cx,bottom,diam=304,605,340
        for value,fill in [(market['tam'],'band'),(market['sam'],'chart_3'),(market['som'],'accent')]:
            dd=diam*math.sqrt(float(value/market['tam']))
            sh.shape(s,MSO_SHAPE.OVAL,*[v/128 for v in (cx-dd/2,bottom-dd,dd,dd)],fill=self.i.color(fill))
        self.i.text(s,170,312,268,72,f'${num(market["tam"],0)}B\nTOTAL ADDRESSABLE',size=26,bold=True,align='c')
        self.i.text(s,218,472,172,58,f'${num(market["sam"])}B\nSERVICEABLE',size=20,bold=True,align='c')
        rows=[(f'${num(market["tam"],0)}B','TAM',market['tam_label']),(f'${num(market["sam"])}B','SAM · '+percent(market['sam']/market['tam']),market['sam_label']),(f'${num(market["som"]*1000,0)}M','FIVE-YEAR SOM',market['som_label'])]
        for j,(value,label,detail) in enumerate(rows):
            yy=236+j*112
            self.i.text(s,640,yy,200,54,value,kind='stat',size=42)
            self.i.text(s,868,yy,356,24,label,kind='label',color='accent')
            self.i.text(s,868,yy+28,356,68,detail,size=18)
            self.i.rule(s,640,yy+102,584)
        self.i.text(s,640,592,584,46,f'{money(b["end"])} ARR today · {percent(b["end"]/(market["som"]*1000),1)} of five-year SOM',size=20,color='muted')

        s,y=self.slide('IN07','Competition','Built for the Whole Quote-to-Order Workflow','Qualitative positioning. Workflow depth and fit for electrical distribution.')
        frame=(144,250,1016,346)
        native=self.ch.scatter_labeled(s,d['competitive_2x2']['players'],x=frame[0],y=frame[1],w=frame[2],h=frame[3],highlight='Parspec',qualitative=True)
        self.plot_layout(native.chart,0,0,1,1)
        self.i.rect(s,*frame,fill=None,line='line')
        self.i.rect(s,frame[0]+frame[2]/2,frame[1],1,frame[3],fill='line')
        self.i.rule(s,frame[0],frame[1]+frame[3]/2,frame[2],color='line')
        self.i.text(s,56,202,1168,25,'BUILT FOR ELECTRICAL DISTRIBUTION (QUALITATIVE)',kind='label',color='muted')
        self.i.text(s,56,250,78,24,'HIGH',kind='label',color='muted',align='r')
        self.i.text(s,56,574,78,24,'LOW',kind='label',color='muted',align='r')
        self.i.text(s,144,620,86,24,'LOW',kind='label',color='muted')
        self.i.text(s,1074,620,86,24,'HIGH',kind='label',color='muted',align='r')
        self.i.text(s,252,620,800,28,'WORKFLOW DEPTH (QUALITATIVE)',kind='label',color='muted',align='c')

        s,y=self.slide('IN08','Valuation','Proposed Valuation / ARR Is Below Peer Median','Valuation / ARR (pre-money). Parspec, Inc. standalone.')
        self.i.text(s,56,240,360,120,ratio(c['multiple']),kind='hero',color='accent')
        self.i.text(s,56,370,360,78,f'Valuation / ARR (pre-money)\nAt ${num(r["pre_money"],0)}M pre-money',size=20,bold=True)
        self.i.text(s,56,466,360,38,'Peer median '+ratio(c['median']),size=20,color='muted')
        self.i.text(s,56,538,360,42,percent(c['growth'],1),size=36,bold=True)
        self.i.text(s,56,590,360,56,'ARR growth, year over year',size=20,color='muted')
        pts=[dict(name=a['name'],x=a['growth'],y=a['valuation_arr']) for a in d['comps']]+[dict(name='Parspec',x=c['growth'],y=c['multiple'])]
        self.ch.scatter_labeled(s,pts,x=448,y=236,w=776,h=366,highlight='Parspec',x_min=.15,x_max=.95,y_min=0,y_max=20,median_x=c['median_growth'],median_y=c['median'],x_format='0%',y_format='0.0"×"')
        self.i.text(s,448,620,776,24,'ARR GROWTH, YEAR OVER YEAR',kind='label',color='muted',align='r')

        s,y=self.slide('IN09','Comparables','Peer Set Detail','Valuation / ARR (pre-money). Parspec, Inc. standalone.')
        self.i.text(s,56,232,330,110,ratio(c['median']),kind='hero',color='accent')
        self.i.text(s,56,354,330,64,'Peer median\nValuation / ARR (pre-money)',size=20,bold=True)
        self.i.rule(s,56,452,310)
        self.i.text(s,56,480,330,84,ratio(c['p25'])+' to '+ratio(c['p75'])+'\nInterquartile range',size=22,bold=True)
        comps=[dict(name=a['name'],growth=a['growth'],multiple=a['valuation_arr']) for a in d['comps']]+[dict(name='Parspec',growth=c['growth'],multiple=c['multiple'])]
        comps.sort(key=lambda a:a['growth'],reverse=True)
        rows=[[a['name'],percent(a['growth'],1 if a['name']=='Parspec' else 0),ratio(a['multiple']),('+' if a['multiple']>=c['median'] else '−')+ratio(abs(a['multiple']-c['median']))] for a in comps]
        self.table(s,['Company','ARR growth','Valuation / ARR\n(pre-money)','Δ vs median'],rows,x=440,y=218,w=784,h=410,widths=[.24,.22,.3,.24],font_px=16,highlight_row=next(j for j,a in enumerate(comps) if a['name']=='Parspec'))
        self.i.text(s,440,632,784,24,'Parspec falls between the lower quartile and the peer median.',size=16,color='muted')

        s,y=self.slide('IN10','The round',f'${num(r["raise"],0)}M Series B at ${num(r["pre_money"],0)}M Pre-Money','Terms and use of proceeds. Parspec, Inc. standalone.')
        self.i.text(s,56,240,400,124,f'${num(r["raise"],0)}M',kind='hero',color='accent')
        self.i.text(s,56,378,400,44,'Series B raise',size=24,bold=True)
        self.i.rule(s,56,448,400)
        for j,(label,v) in enumerate([('Pre-money',r['pre_money']),('Post-money',r['pre_money']+r['raise'])]):
            self.i.text(s,56+j*220,478,218,64,money(v,0),kind='stat',size=42)
            self.i.text(s,56+j*220,548,218,30,label,size=18,color='muted')
        self.i.text(s,56,612,400,30,percent(r['raise']/(r['pre_money']+r['raise']))+' new-money share',size=20,color='muted')
        self.ch.doughnut(s,[a['use'] for a in r['use_of_funds']],[a['pct'] for a in r['use_of_funds']],x=510,y=244,w=412,h=348)
        self.i.text(s,636,392,160,50,money(r['raise'],0),size=38,bold=True,align='c')
        for j,a in enumerate(r['use_of_funds']):
            yy=264+j*86;self.i.rect(s,956,yy+7,16,16,fill=f'chart_{j+1}')
            self.i.text(s,988,yy,236,28,a['use'],size=18,bold=True)
            self.i.text(s,988,yy+34,236,28,money(r['raise']*a['pct'])+' · '+percent(a['pct']),size=16,color='muted')

        s,y=self.slide('IN11','Plan','Three Milestones This Round Funds','Today: Q3 FY26 '+money(b['end'])+' ARR. Parspec, Inc. standalone.')
        self.i.rule(s,92,378,1090,color='accent',h=2)
        for j,m in enumerate(r['milestones']):
            xx=132+j*374
            sh.shape(s,MSO_SHAPE.OVAL,*[v/128 for v in (xx,365,26,26)],fill=self.i.color('ground'),line=self.i.color('accent'),lw=1)
            self.i.text(s,xx-76,298,220,40,f'{j+1:02}',size=20,bold=True,color='accent',align='c')
            self.i.text(s,xx-76,420,330,104,m['what'].replace(' ending ARR','\nEnding ARR').replace(' run rate','\nRun rate'),size=40,bold=True,align='c',color='accent' if j==2 else 'text')
            self.i.text(s,xx-76,550,330,32,m['by'],size=20,color='muted',align='c')

        s,y=self.slide('IN12','Appendix','Sources and Definitions','Illustrative sample data (template fixture v4). Parspec, Inc. standalone.')
        rows=[['ARR','Ending recurring revenue. Actual through Q3 FY26; forecast from Q4 FY26.','arr_history; arr_bridge_q3_fy26'],['Company NDR / GRR','Trailing-twelve-month fixture assertions, separate from the cohort table.','unit_economics; kpi_note'],['Cohort expansion','Cumulative ARR relative to Year 1. Blank cells are unavailable.','cohort_ndr'],['Valuation / ARR (pre-money)','Pre-money equity value / ending ARR for company and peers. This is not EV.','valuation_basis; comps; round'],['Peer statistics','Unrounded median and inclusive interpolated quartiles; deltas precede display rounding.','comps'],['Cash / runway / cases','Twelve months of cash; runway and later cash-out are supplied assertions. Minimum-cash horizon unspecified.','cash_runway; cases; runway_note'],['Market / unit economics','Illustrative assumptions. SOM is a five-year opportunity. CAC / LTV are assertions.','market; unit_economics'],['Rounding / fiscal basis','Store two decimals half-up; sum stored quarters. Fiscal year opens in February. FY26 ends January 2027.','rounding_rule; fiscal_year'],['Customer / team / product','Fictional asserted outcomes and annual benefits; credentials and attribution remain placeholders. Track orders is planned; other listed steps live.','customer_outcomes; customer_value; team; product_mechanism'],['EBITDA to cash','Same-quarter supplied adjustments sum to negative net burn; this is an illustrative reconciliation.','ebitda_to_cash_q3_fy26']]
        self.table(s,['Metric','Definition / limitation','Fixture fields'],rows,y=202,h=446,widths=[.2,.55,.25],font_px=14,numeric_cols=[])

        s,y=self.slide('IN13','Important information','Confidential; for Discussion Only','Parspec, Inc. standalone.')
        paragraphs=[('FICTIONAL TEMPLATE','This version uses fictional sample data. It is an illustrative design template and does not report Parspec financial results or substantiate any investment claim.'),('CONFIDENTIALITY','This presentation is confidential and for discussion only. It is not an offer to sell or a solicitation to buy any security. Any investment requires definitive agreements.'),('FORECASTS AND METRICS','Plans are forward-looking estimates. Actual results may differ materially. ARR, retention and CAC payback are operating metrics. Definitions and fixture limitations appear on the Sources and Definitions page.'),('VALUATION AND SOURCES','Peer and market values in this template are fictional assumptions. Valuation references use pre-money equity value / ARR. They are not a fairness opinion, independent appraisal or tax valuation.')]
        for j,(label,body) in enumerate(paragraphs):
            yy=206+j*108
            self.i.text(s,56,yy,1168,25,label,kind='label',color='accent')
            self.i.text(s,56,yy+31,1168,68,body,size=18,color='muted')

        s,y=self.slide('IN14','Financial plan',f'Path to ${num(c["arr"][-1],0)}M ARR in FY27','FY26 total: actual + forecast. Parspec, Inc. standalone.')
        self.ch.act_forecast_chips(s,'ACTUAL · THROUGH Q3 FY26','FORECAST · FROM Q4 FY26',x=450,y=184,w=774)
        self.ch.column_actual_forecast(s,d['arr_history']['quarters'][6:],c['arr'][6:],3,x=360,y=226,w=864,h=180)
        p=d['pnl']; fy=d['fy27_plan']; rev=p['revenue']+fy['revenue']; gp=c['gp']+c['gp27']; nn=c['net_new']
        rows=[['Ending ARR']+[num(v) for v in c['arr'][6:]]+[num(c['arr'][9]),num(c['arr'][-1])],['Net new ARR']+[num(v) for v in nn]+[num(sum(nn[:4])),num(sum(nn[4:]))],['Revenue']+[num(v) for v in rev]+[num(sum(rev[:4])),num(sum(rev[4:]))],['Gross profit']+[num(v) for v in gp]+[num(sum(gp[:4])),num(sum(gp[4:]))],['Gross margin']+[percent(g/r) for g,r in zip(gp,rev)]+[percent(sum(gp[:4])/sum(rev[:4])),percent(sum(gp[4:])/sum(rev[4:]))]]
        self.table(s,['US$M']+p['quarters']+fy['quarters']+['FY26','FY27'],rows,y=434,h=212,widths=[.19]+[.075]*8+[.105,.105],font_px=14,statuses=['']+['Actual']*3+['Forecast']*5+['Mixed','Forecast'])
        self.customer_outcomes()
        self.customer_value()
        self.team()
        self.product_mechanism()
        self.cash_bridge('IN19')
        return self.deck


def compare(out,picks_root,render_dirs):
    from PIL import Image, ImageDraw
    records=[]
    for k,(sid,lane) in enumerate(PICKS):
        board=sid.startswith('BD'); index=(BOARD_IDS if board else INVESTOR_IDS).index(sid)+1
        stem='parspec_board_gallery' if board else 'parspec_investor_gallery'
        root=render_dirs[0 if board else 1]
        candidates=list(root.glob(f'{stem}-*.png'))
        native=next(p for p in candidates if int(p.stem.rsplit('-',1)[1])==index)
        picked=picks_root/lane/('01.png' if sid in ('BD01','IN01') else sid+'.png')
        a=Image.open(picked).convert('RGB').resize((1280,720)); b=Image.open(native).convert('RGB').resize((1280,720))
        canvas=Image.new('RGB',(2560,752),'white');canvas.paste(a,(0,32));canvas.paste(b,(1280,32))
        draw=ImageDraw.Draw(canvas);draw.text((16,8),f'{sid} PICK / {lane.upper()}',fill='black');draw.text((1296,8),f'{sid} NATIVE / v4 fixture',fill='black')
        target=out/f'compare_{sid[:2]}-{sid[2:]}.png';canvas.save(target)
        records.append(dict(slide=sid,pick=lane,comparison=str(target),native_sha256=hashlib.sha256(native.read_bytes()).hexdigest(),picked_sha256=hashlib.sha256(picked.read_bytes()).hexdigest()))
    (out/'comparisons.json').write_text(json.dumps(records,indent=2)+'\n')
    new_records=[]
    import shutil
    for sid in ['BD16']+[f'IN{i:02}' for i in range(15,20)]:
        board=sid.startswith('BD'); index=(BOARD_IDS if board else INVESTOR_IDS).index(sid)+1
        stem='parspec_board_gallery' if board else 'parspec_investor_gallery'
        root=render_dirs[0 if board else 1]
        native=next(p for p in root.glob(f'{stem}-*.png') if int(p.stem.rsplit('-',1)[1])==index)
        target=out/f'render_{sid[:2]}-{sid[2:]}.png';shutil.copy2(native,target)
        new_records.append(dict(slide=sid,render=str(target),sha256=hashlib.sha256(native.read_bytes()).hexdigest()))
    (out/'new_slide_renders.json').write_text(json.dumps(new_records,indent=2)+'\n')


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--data',type=Path,required=True)
    ap.add_argument('--out',type=Path,default=TEMPLATE_DIR)
    ap.add_argument('--render',action='store_true')
    ap.add_argument('--evidence',type=Path)
    ap.add_argument('--picks-root',type=Path)
    args=ap.parse_args()
    raw=args.data.read_bytes();d=json.loads(raw,parse_float=D)
    if 'FICTIONAL' not in d.get('_note','') or 'v4' not in d.get('_note',''):
        ap.error('requires the explicitly fictional v4 fixture')
    args.out.mkdir(parents=True,exist_ok=True)
    sha=hashlib.sha256(raw).hexdigest();render_dirs=[]; manifests=[]
    for name,build in [('board',Gallery.board),('investor',Gallery.investor)]:
        g=Gallery(d,sha);build(g)
        for slide in g.deck.prs.slides:
            # Cover artwork bleeds behind the common footer, never over it.
            for shape in list(slide.shapes):
                if shape.top / 914400 * 128 >= 680:
                    slide.shapes._spTree.remove(shape._element)
                    slide.shapes._spTree.append(shape._element)
        path=args.out/f'parspec_{name}_gallery.pptx';g.deck.save(path)
        manifests.append(dict(deck=str(path),slides=g.ids,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),fixture_sha256=sha))
        print(f'{path}: {len(g.ids)} slides')
        root=args.out/f'{name}_gallery_png';render_dirs.append(root)
        if args.render:
            env=dict(os.environ);env.setdefault('SOFFICE_PROFILE','file:///tmp/sol_lo_profile')
            subprocess.run(['bash',str(SKILL/'scripts/render.sh'),str(path),str(root),'128'],env=env,check=True)
    if args.evidence:
        args.evidence.mkdir(parents=True,exist_ok=True)
        (args.evidence/'gallery_manifest.json').write_text(json.dumps(manifests,indent=2)+'\n')
        if args.render and args.picks_root:compare(args.evidence,args.picks_root,render_dirs)


if __name__=='__main__':main()
