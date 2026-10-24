"""Independent story-deck oracle; saved OOXML, embedded workbooks and fixtures."""
from copy import deepcopy
from decimal import Decimal as D
import hashlib
import json
import re
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.util import Pt
from investor_smoke import require, frames, workbook_contract, values_equal, staging_contract, FONT

IDS=['IN-20','IN-21','IN-22','INCUBATION']+[f'IN-{i}' for i in range(23,32)]
CHARTS={'IN-23':1,'IN-24':1,'IN-27':1,'IN-28':2,'IN-31':3}
CHIP='ILLUSTRATIVE · DUMMY NUMBERS'
FOOTER='Illustrative — dummy numbers, not Parspec data'


def matrix(slide):
    return [[[c.text for c in r.cells] for r in t.table.rows] for t in slide.shapes if t.has_table]


def text(slide):
    return '\n'.join(f.text for _,f in frames(slide))


def chart_values(chart):
    return [list(s.values) for s in chart.series]


def floating(steps,totals):
    base=[];movement=[];value=0
    for j,v in enumerate(steps):
        if j in totals:base.append(0);movement.append(v);value=v
        else:base.append(min(value,value+v));movement.append(abs(v));value+=v
    return base,movement


def check_slide(slide,sid,d):
    t=text(slide);inc=sid not in ('IN-20','IN-21','IN-22')
    require(t.count(CHIP)==int(inc),sid+': illustrative boundary')
    require((FOOTER in t)==inc,sid+': incubation footer')
    if not inc:
        require(not re.search(r'\b(?:UNKNOWN|TBD)\b',t),sid+': raw machine placeholder on real-content slide')
    footer_texts=[f.text for shape,f in frames(slide) if shape.top/914400*128>=680]
    require(any((FOOTER if inc else '©2026 Parspec Inc. All rights reserved. Proprietary and confidential.')==value for value in footer_texts),sid+': complete footer')
    for shape,frame in frames(slide):
        floor=14
        for p in frame.paragraphs:
            for r in p.runs:
                if not r.text:continue
                require(r.font.size is not None,sid+': inherited font')
                require(r.font.size.pt*128/72+.02>=floor,sid+': type floor')
                require(r.font.name==FONT,sid+': font family')
        require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=9144000+100 and shape.top+shape.height<=5143500+100,sid+': text off canvas')
    charts=[s.chart for s in slide.shapes if s.has_chart]
    require(len(charts)==CHARTS.get(sid,0),sid+': native chart count')
    if sid not in ('INCUBATION','IN-21','IN-28'):require(matrix(slide),sid+': native table missing')
    for ch in charts:
        workbook_contract(ch,sid)
        for axis in ch._chartSpace.xpath('.//c:valAx/c:scaling/c:min'):
            require(float(axis.get('val'))==0,sid+': zero axis required')
        for sz in ch._chartSpace.xpath('.//a:rPr/@sz | .//a:defRPr/@sz'):
            require(float(sz)/100*128/72+.02>=14,sid+': chart text floor')
    if sid=='INCUBATION':
        require('Evidence needed to test the economics' in t.replace('\n',' '),'divider: test, not promised proof')
        require('Independent illustrative examples, not a connected financial model' in t,'divider: independent examples')
    if sid=='IN-20':
        require(len(slide._element.xpath('.//p:cxnSp'))>=2,'IN20: connectors')
    if sid=='IN-21':
        require(len(slide._element.xpath('.//p:cxnSp'))>=2,'IN21: connectors')
        require(not any(sh.name.startswith('control:') and sh._element.xpath('.//a:tailEnd[@type="triangle"]') for sh in slide.shapes),'IN21: human controls must not form a chronological arrow chain')
        feeds=[sh for sh in slide.shapes if sh.name=='intake:feed']
        require(feeds,'IN21: alternative triggers must feed agent work')
        # A populated routed draft carries its claim limits visibly. Generic
        # gallery steps remain explicit placeholders with example statuses.
        if any(sh.name.startswith('route:') for sh in slide.shapes):
            status=next(sh.text for sh in slide.shapes if sh.name=='flow:status-1')
            require(status=='Design specified','B01: extraction execution is not step-verified')
            require('Quote-draft workflow with human review' in t.replace('\n',' '),'B01: narrowed headline')
            require('not in customer production' in t,'B01: maturity boundary')
            require('PROPOSED\nHUMAN CONTROLS' in t,'B01: proposed controls')
            require('Project RFQ' in t and 'Stock order' in t,'F03: branch labels')
    if sid in ('IN-23','IN-29','IN-30','IN-31'):
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                require(not run._r.xpath('./a:rPr[@spc>0]'),sid+': tracked numeric glyphs can split numbers')
    if sid=='IN-23':
        rows=d['IN23']['rows'];total=sum(D(r['parents'])*(D(str(r['platform_fee']))+D(r['units_per_parent'])*D(str(r['net_unit_price']))) for r in rows)/1000000
        require(total.quantize(D('.1'))==D('301.2'),'IN23: operand total')
        values_equal(chart_values(charts[0])[0],[total,d['IN23']['expansion_scenario_m']],'IN23 values')
        require(matrix(slide)[0][-1][-1]=='301.2','IN23: displayed total')
        require(charts[0]._chartSpace.xpath('.//c:barChart/c:grouping[@val="clustered"]'),'IN23: clustered columns')
    if sid=='IN-24':
        steps=d['IN24']['steps'];running=0
        for r in steps:
            if r['type']=='total':
                if running:require(r['value']==running,'IN24 total endpoint')
                running=r['value']
            else:running+=r['value']
        require(steps[0]['value']==4000 and running==300,'IN24: 4000 to 300')
        expected=floating([r['value'] for r in steps],{j for j,r in enumerate(steps) if r['type']=='total'})
        for got,want in zip(chart_values(charts[0])[:2],expected):values_equal(got,want,'IN24 waterfall')
        require(charts[0].series[0]._element.xpath('./c:spPr/a:noFill'),'IN24 hidden base')
        cats=[c.label.replace('\n',' ') for c in charts[0].plots[0].categories]
        require('Outside technical coverage (40%)' in cats and 'Unrealized benefit (30%)' in cats,'F06: removed shares')
        require('60% retained: 4,000 → 2,400' in t and '70% realized: 2,400 → 1,680' in t,'F06: retained-factor table')
    if sid=='IN-25':
        tables=matrix(slide);require(len(tables)==2,'IN25 marketplace separation')
        require('Marketplace' in tables[1][1][0] and 'EXCLUDED FROM EVERY TOTAL' in t,'IN25 marketplace excluded')
    if sid=='IN-26':
        require(any(f.text.startswith('Proposed pilot package') for _,f in frames(slide)),'IN26 title')
        require('Meter / billing basis' in t,'F07: meter versus billing')
        require('Committed credits billed in advance; documents consume credits' in t and 'Committed D1 project-months billed in advance' in t,'F07: advance billing')
    if sid=='IN-27':
        require('US$ / accepted output · independent of the pilot-package example' in t,'F05: independent economics')
        p=d['IN27']['per_accepted_output'];ch=charts[0]
        require(ch._chartSpace.xpath('.//c:barChart/c:grouping[@val="clustered"]'),'IN27 columns must be adjacent')
        require(len(ch.series)==1,'IN27: single series, three categories')
        values_equal(chart_values(ch)[0],[p['customer_benefit'],p['net_price'],p['full_cost']],'IN27 values')
        require(len(slide._element.xpath('.//p:cxnSp'))>=6,'IN27 editable brackets')
        require('$15.00' in t and '$1.80' in t,'IN27 bracket differences')
    if sid=='IN-28':
        require('US$K · waterfall period not specified · payback is a separate illustration' in t,'F08: unspecified period and separate illustration')
        rows=d['IN28']['waterfall_k'];gp=sum(r['value'] for r in rows[:-1]);require(gp==522==rows[-1]['value'],'IN28 GP')
        expected=floating([r['value'] for r in rows],{0,len(rows)-1})
        for got,want in zip(chart_values(charts[0])[:2],expected):values_equal(got,want,'IN28 waterfall')
        p=d['IN28']['payback'];cum=[];v=0
        for n in p['monthly_gp_k']:v+=n;cum.append(v)
        month=next(j+1 for j,n in enumerate(cum) if n>=p['deploy_cost_k'])
        require(month==13==p['payback_month'],'IN28 payback month')
        values_equal(chart_values(charts[1])[0],cum,'IN28 cumulative GP')
        values_equal(chart_values(charts[1])[1],[p['deploy_cost_k']]*len(cum),'IN28 deployment cost')
        require('month 13' in t and charts[1]._chartSpace.xpath('.//a:t[text()="Month 13"]'),'IN28 payback marked')
    if sid=='IN-29':
        require('All signed customers are shown. Later-cohort renewal status was not provided.' in t,'F09: no inferred maturity')
        for row,source in zip(matrix(slide)[0][1:],d['IN29']):
            if source['renewed'] is None:require(row[4]=='Not provided','F09: null is not zero or unmatured')
    if sid=='IN-30':
        require(len(matrix(slide)[0])==len(d['IN30'])+1,'IN30 evidence coverage')
        require('Baseline rate' in t and 'With-context rate' in t,'F10: rate headers')
        require(t.count('500 RFQs; metric to be defined')==2,'F10: missing success definition')
        require('Illustrative rates; success definitions and evaluation rules remain to be specified' in t,'F10: rates are illustrative')
    if sid=='IN-31':
        ch=charts[0];require(ch._chartSpace.xpath('.//c:barChart/c:grouping[@val="percentStacked"]'),'IN31 100% stacked')
        values=chart_values(ch)
        for j,row in enumerate(d['IN31']):
            require(sum(row[k] for k in ('accepted_no_review_pct','accepted_with_review_pct','failed_pct'))==100,'IN31 fixture sums')
            require(abs(sum(v[j] for v in values)-1)<1e-8,'IN31 chart sums')
        for j,key in enumerate(('accepted_no_review_pct','accepted_with_review_pct','failed_pct')):
            values_equal(values[j],[D(r[key])/100 for r in d['IN31']],'IN31 outcomes')
        for ch,key in zip(charts[1:],('review_min_per_output','full_cost_per_output')):
            values_equal(chart_values(ch)[0],[r[key] for r in d['IN31']],'IN31 line')
        require('TBD' in t,'IN31 unsupported period / defects stay unknown')


def probes(prs,d):
    results=[]
    def reject(sid,name,mutate):
        source=prs.slides[IDS.index(sid)];s=Slide(deepcopy(source._element),source.part)
        restore=mutate(s)
        try:
            try:check_slide(s,sid,d)
            except AssertionError:results.append(name)
            else:raise AssertionError('mutation survived: '+name)
        finally:
            if restore:restore()
    def remove_chip(s):
        shape=next(s for s in s.shapes if s.has_text_frame and s.text==CHIP);shape._element.getparent().remove(shape._element)
    reject('IN-23','missing illustrative chip',remove_chip)
    def wrong_boundary(s):
        chip=next(sh for sh in prs.slides[4].shapes if sh.has_text_frame and sh.text==CHIP)
        s.shapes._spTree.append(deepcopy(chip._element))
    reject('IN-20','dummy chip on real-content slide',wrong_boundary)
    def small(s):
        shape=next(s for s in s.shapes if s.has_text_frame and s.text==CHIP);shape.text_frame.paragraphs[0].runs[0].font.size=Pt(5)
    reject('IN-23','font floor',small)
    def footer_small(s):
        box=next(sh for sh in s.shapes if sh.has_text_frame and sh.text==FOOTER)
        box.text_frame.paragraphs[0].runs[0].font.size=Pt(6.19)
    reject('IN-23','footer below 14px',footer_small)
    for literal in ('UNKNOWN','TBD'):
        def raw_placeholder(s,value=literal):
            box=next(sh for sh in s.shapes if sh.has_text_frame and sh.top/914400*128<680)
            box.text_frame.paragraphs[0].runs[0].text=value
        reject('IN-20','raw '+literal+' on real-content slide',raw_placeholder)

    def tracked(s):
        table=next(sh.table for sh in s.shapes if sh.has_table)
        run=table.cell(1,1).text_frame.paragraphs[0].runs[0]
        run._r.get_or_add_rPr().set('spc','200')
    reject('IN-23','tracked digits break numeric legibility',tracked)
    if any(sh.name.startswith('route:') for sh in prs.slides[1].shapes):
        def strengthen_extraction(s):
            box=next(sh for sh in s.shapes if sh.name=='flow:status-1')
            box.text_frame.paragraphs[0].runs[0].text='Demo environment'
        reject('IN-21','B01 unsupported extraction execution',strengthen_extraction)
    def chronological_controls(s):
        sh=next(sh for sh in s.shapes if sh.name=='control:association')
        from pptx.oxml.xmlchemy import OxmlElement
        arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');sh._element.spPr.get_or_add_ln().append(arrow)
    reject('IN-21','F03 chronological human controls',chronological_controls)

    def stacked(s):
        ch=next(s.chart for s in s.shapes if s.has_chart);g=ch._chartSpace.xpath('.//c:barChart/c:grouping')[0];old=g.get('val');g.set('val','stacked')
        return lambda:g.set('val',old)
    reject('IN-27','stacked overlapping economics',stacked)
    for sid in ('IN-24','IN-28','IN-31'):
        def corrupt(s):
            ch=next(s.chart for s in s.shapes if s.has_chart);v=ch._chartSpace.xpath('.//c:numCache/c:pt/c:v')[0];old=v.text;v.text='99999'
            return lambda:setattr(v,'text',old)
        reject(sid,sid+' corrupt chart cache',corrupt)
    return results


def run(data,paths,evidence=None):
    raw=Path(data).read_bytes();d=json.loads(raw);sha=hashlib.sha256(raw).hexdigest();decks=[];mutations=[]
    for path in paths:
        prs=Presentation(path);require(len(prs.slides)==13,'story count')
        for sid,s in zip(IDS,prs.slides):
            require('Slide: '+sid in s.notes_slide.notes_text_frame.text,sid+': ID/order')
            if sid.startswith('IN-') and int(sid[3:])>=23:require(sha in s.notes_slide.notes_text_frame.text,sid+': fixture provenance')
            check_slide(s,sid,d)
        mutations+=probes(prs,d)
        decks.append(dict(path=str(path),slides=13,native_charts=sum(sh.has_chart for s in prs.slides for sh in s.shapes),sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest()))
    staged=staging_contract()
    receipt=dict(result='PASS',decks=decks,fixture_sha256=sha,rejected_mutations=mutations,staged_paths_checked=staged,limitations='Structural/data check. Render inspection and external readiness are separate.')
    if evidence:Path(evidence).write_text(json.dumps(receipt,indent=2)+'\n')
    print(f'PASS: {len(decks)*13} story slides; {sum(d["native_charts"] for d in decks)} native charts; {len(mutations)} rejected mutations; fixture arithmetic, workbooks, incubation boundary and staging guard')
    return receipt
