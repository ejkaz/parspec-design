"""Independently verify local fictional investor / board gallery artifacts.

The oracle reads the caller's v3 fixture, never the gallery manifest or builder
calculations. It checks native chart caches against independently derived data,
their embedded workbooks, financial-table matrices, per-slide number provenance,
type floors, fiscal labels and public-repository staging. Mutation probes run by
default. This is a structural / data check, not visual or brand acceptance.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from pptx import Presentation
from pptx.oxml.ns import qn

SKILL = Path(__file__).resolve().parents[1]
REPO = SKILL.parents[3]
sys.path.insert(0, str(SKILL))
from assets.brand import load as load_brand
from assets.house import DEFAULT_PROFILE, TEMPLATE_DIR

D = Decimal
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
BOARD_IDS = ['BD01', 'BD01a'] + [f'BD{i:02}' for i in range(2, 16)]
INVESTOR_IDS = [f'IN{i:02}' for i in range(1, 15)]
CHART_SLIDES = {'BD03', 'BD04', 'BD07', 'BD08', 'BD09', 'BD10',
                'IN03', 'IN05', 'IN07', 'IN08', 'IN10', 'IN14'}
TABLE_SLIDES = {'BD01a', 'BD02', 'BD05', 'BD06', 'BD07', 'BD08',
                'BD09', 'BD10', 'BD12', 'BD13', 'BD14', 'IN04',
                'IN09', 'IN12', 'IN14'}
NUMBER = re.compile(r'(?<![A-Za-z0-9_])\d+(?:,\d{3})*(?:\.\d+)?')
MONTH = re.compile(r'\b(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+20\d{2}\b', re.I)
FONT = load_brand().font


def require(ok, message):
    if not ok:
        raise AssertionError(message)


def rounded(value, places=2):
    return D(str(value)).quantize(D(1).scaleb(-places), rounding=ROUND_HALF_UP)


def fmt(value, places=1):
    if value is None:
        return '—'
    v = rounded(value, places)
    return f'({abs(v):.{places}f})' if v < 0 else f'{v:.{places}f}'


def pct(value, places=0):
    return fmt(value * 100, places) + '%'


def multiple(value):
    return fmt(value) + '×'


def independent(d):
    """Financial operands, Decimal rounding and inclusive statistics only."""
    p, f = d['pnl'], d['fy27_plan']
    gp = [rounded(rev * margin) for rev, margin in zip(p['revenue'], p['gross_margin'])]
    gp27 = [rounded(rev * margin) for rev, margin in zip(f['revenue'], f['gross_margin'])]
    ebitda = [rounded(g - o) for g, o in zip(gp, p['opex'])]
    cor = [rev - g for rev, g in zip(p['revenue'], gp)]
    weights = [rounded(row['amount'] * row['win_rate']) for row in d['pipeline_q4_fy26']['stages']]
    arr = d['arr_history']['arr']
    peer = sorted(row['valuation_arr'] for row in d['comps'])
    growths = sorted(row['growth'] for row in d['comps'])
    def quantile(frac):
        rank = D(len(peer) - 1) * frac
        lo = int(rank)
        return peer[lo] + (rank - lo) * (peer[min(lo + 1, len(peer) - 1)] - peer[lo])
    return dict(gp=gp, gp27=gp27, ebitda=ebitda, cor=cor, weights=weights,
                weighted=sum(weights), coverage=sum(weights) / d['pipeline_q4_fy26']['target_net_new'],
                median=(peer[3] + peer[4]) / 2, growth_median=(growths[3] + growths[4]) / 2,
                p25=quantile(D('.25')), p75=quantile(D('.75')),
                multiple=d['round']['pre_money'] / d['arr_bridge_q3_fy26']['end'],
                yoy=arr[8] / arr[4] - 1, two_year=arr[8] / arr[0],
                net_new=[arr[i] - arr[i - 1] for i in range(6, len(arr))])


def numeric_allowances(d, c):
    """Per-page source branches plus named computations, never global values."""
    branch = {
        'BD01': [], 'BD01a': ['agenda', 'decisions'],
        'BD02': ['arr_bridge_q3_fy26', 'unit_economics', 'pnl', 'plan_vs_actual_q3_fy26', 'quarter_kpis'],
        'BD03': ['arr_bridge_q3_fy26', 'unit_economics'],
        'BD04': ['arr_history'], 'BD05': ['plan_vs_actual_q3_fy26'],
        'BD06': ['pnl'], 'BD07': ['cash_runway', 'plan_vs_actual_q3_fy26'],
        'BD08': ['cases'], 'BD09': ['headcount'],
        'BD10': ['pipeline_q4_fy26'], 'BD11': ['decisions'],
        'BD12': ['pnl'], 'BD13': ['cash_runway', 'decisions'],
        'BD14': ['risks'], 'BD15': ['agenda'],
        'IN01': [], 'IN02': ['arr_history', 'unit_economics', 'pnl', 'market'],
        'IN03': ['arr_history'], 'IN04': ['cohort_ndr', 'unit_economics'],
        'IN05': ['unit_economics'], 'IN06': ['market', 'arr_bridge_q3_fy26'],
        'IN07': ['competitive_2x2'], 'IN08': ['comps', 'round', 'arr_history'],
        'IN09': ['comps', 'round', 'arr_history'], 'IN10': ['round'],
        'IN11': ['round', 'arr_bridge_q3_fy26'], 'IN12': [],
        'IN13': [], 'IN14': ['arr_history', 'pnl', 'fy27_plan'],
    }
    allowed = {sid: set() for sid in branch}
    def add(sid, value, as_rate=False):
        if isinstance(value, dict):
            for key, item in value.items(): add(sid, item, as_rate or key in {'ndr', 'grr', 'gm', 'gross_margin', 'growth', 'win_rate', 'pct'})
        elif isinstance(value, (list, tuple)):
            for item in value: add(sid, item, as_rate)
        elif isinstance(value, (int, float, D)) and not isinstance(value, bool):
            v = abs(D(str(value)))
            for places in (0, 1, 2, 3): allowed[sid].add(abs(rounded(v, places)))
            allowed[sid].add(v)
            if as_rate or 0 <= v <= 1:
                for places in (0, 1, 2): allowed[sid].add(abs(rounded(v * 100, places)))
        elif isinstance(value, str):
            for match in NUMBER.finditer(value): allowed[sid].add(D(match.group().replace(',', '')))
    for sid, fields in branch.items():
        # Page indices, sequence numbers, fiscal quarter numbers, explicit
        # calendar years and Year-1 / two-year references are structural.
        add(sid, [1, 2, 3, 4, 2026, 2027])
        add(sid, BOARD_IDS.index(sid) + 1 if sid in BOARD_IDS else INVESTOR_IDS.index(sid) + 1)
        for field in fields: add(sid, d[field])
    b, cash, hc = d['arr_bridge_q3_fy26'], d['cash_runway'], d['headcount']
    add('BD01a', [sum(x['minutes'] for x in d['agenda']), list(range(1, len(d['agenda']) + 1))])
    for sid in ('BD02', 'BD03'):
        add(sid, [b['end'] - b['start'], b['new'] + b['expansion'], b['contraction'] + b['churn'],
                  d['plan_vs_actual_q3_fy26'][0]['plan'] - b['start']])
    add('BD04', [c['yoy'], (d['arr_history']['arr'][-1] / d['arr_history']['arr'][9] - 1),
                 d['arr_history']['arr'][-1] - d['arr_history']['arr'][9]])
    add('BD05', [r['actual'] - r['plan'] for r in d['plan_vs_actual_q3_fy26']])
    for sid in ('BD06', 'BD12'):
        add(sid, [c['gp'], c['ebitda'], c['cor'], sum(d['pnl']['revenue']), sum(c['gp']),
                  sum(d['pnl']['opex']), sum(c['ebitda']), sum(c['cor']),
                  d['pnl']['revenue'][-1] - d['pnl']['revenue'][0]])
        add(sid, [g / r for g, r in zip(c['gp'], d['pnl']['revenue'])])
        add(sid, [e / r for e, r in zip(c['ebitda'], d['pnl']['revenue'])], as_rate=True)
        add(sid, [sum(c['gp']) / sum(d['pnl']['revenue']), sum(c['ebitda']) / sum(d['pnl']['revenue'])], as_rate=True)
        add(sid, [d['pnl']['revenue'][j] / d['pnl']['revenue'][j - 1] - 1 for j in range(1, 4)])
    for sid in ('BD07', 'BD13'):
        add(sid, [sum(cash['net_burn']), sum(cash['net_burn']) / len(cash['net_burn']),
                  [v - cash['debt_draw']['amount'] for v in cash['cash_end'][3:]]])
    for field in ('start', 'joiners', 'leavers', 'end', 'plan_q4'): add('BD09', sum(hc[field]))
    add('BD09', [sum(hc['end']) - sum(hc['start']), sum(hc['plan_q4']) - sum(hc['end'])])
    add('BD10', [c['weights'], c['weighted'], c['coverage'], sum(c['weights'][2:]),
                 sum(s['amount'] for s in d['pipeline_q4_fy26']['stages'])])
    add('IN02', [c['two_year'], (d['unit_economics']['gross_margin_trend']['gm'][-1] - d['unit_economics']['gross_margin_trend']['gm'][0]) * 100,
                 d['market']['som'] * 1000])
    add('IN03', c['two_year'])
    add('IN05', (d['unit_economics']['gross_margin_trend']['gm'][-1] - d['unit_economics']['gross_margin_trend']['gm'][0]) * 100)
    m = d['market']
    add('IN06', [m['som'] * 1000, m['sam'] / m['tam'], b['end'] / (m['som'] * 1000)])
    for sid in ('IN08', 'IN09'):
        add(sid, [c['multiple'], c['median'], c['growth_median'], c['p25'], c['p75'], c['yoy'],
                  [row['valuation_arr'] - c['median'] for row in d['comps']], c['multiple'] - c['median']])
    r = d['round']
    add('IN10', [r['pre_money'] + r['raise'], r['raise'] / (r['pre_money'] + r['raise']),
                 [r['raise'] * use['pct'] for use in r['use_of_funds']]])
    add('IN14', [c['net_new'], sum(c['net_new'][:4]), sum(c['net_new'][4:]), c['gp'], c['gp27'],
                 sum(c['gp']), sum(c['gp27']), sum(d['pnl']['revenue']), sum(d['fy27_plan']['revenue'])])
    revenues = d['pnl']['revenue'] + d['fy27_plan']['revenue']
    profits = c['gp'] + c['gp27']
    add('IN14', [g / rev for g, rev in zip(profits, revenues)])
    add('IN14', [sum(c['gp']) / sum(d['pnl']['revenue']), sum(c['gp27']) / sum(d['fy27_plan']['revenue'])])
    return allowed


def frames(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape, shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield shape, cell.text_frame


def fiscal_contract(text, sid, nearby=''):
    for match in re.finditer(r'\bQ[1-4]\b', text):
        require(re.match(r'\s+FY\d{2,4}\b', text[match.end():]), f'{sid}: bare quarter in {text!r}')
    for match in MONTH.finditer(text):
        # Definitions may state the fiscal-year end in one sentence. Main
        # calendar-month labels must carry a fiscal-quarter label in that box.
        month_name, year_text = match.group().split()
        month = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'].index(month_name[:3].lower()) + 1
        year = int(year_text); fy = year if month >= 2 else year - 1
        quarter = ((month - 2) % 12) // 3 + 1
        expected = f'Q{quarter} FY{str(fy)[-2:]}'
        require(sid in {'BD12', 'IN12'} or expected in text + ' ' + nearby,
                f'{sid}: calendar month lacks fiscal quarter in {text!r}')


def text_contract(slide, sid, allowed):
    footer_found = False
    for shape, frame in frames(slide):
        text = frame.text
        if not text.strip(): continue
        footer = shape.top / 914400 * 128 >= 672
        footer_found |= footer and '©' in text
        floor = 11 if footer else 13 if sid in {'BD12', 'IN09'} else 14
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if not run.text: continue
                size = run.font.size or paragraph.font.size
                require(size is not None, f'{sid}: inherited font size is not measurable: {run.text!r}')
                require(run.font.name == FONT, f'{sid}: nonbrand text font {run.font.name!r}')
                require(size.pt * 128 / 72 + .02 >= floor,
                        f'{sid}: {size.pt * 128 / 72:.2f}px below {floor}px: {run.text!r}')
                if sid not in {'BD12', 'IN12'}:
                    require(not run._r.xpath('./a:rPr[@baseline>0]'), f'{sid}: superscript run')
        if not footer:
            # A tile may split its calendar value and fiscal context across
            # adjacent text boxes. Require the associated label in that region,
            # rather than accepting an unrelated fiscal label elsewhere.
            nearby = ' '.join(other.text for other in slide.shapes
                              if other.has_text_frame and other is not shape
                              and abs(other.top - shape.top) / 914400 * 128 <= 130
                              and other.left < shape.left + shape.width
                              and shape.left < other.left + other.width)
            fiscal_contract(text, sid, nearby)
            if sid not in {'BD12', 'IN12'}:
                require(not re.search(r'\bn\s*=|\bfootnote\b|\bmethodology\b|\bsource\s*:', text, re.I),
                        f'{sid}: main-slide footnote / methodology text')
        for token in NUMBER.findall(text):
            require(D(token.replace(',', '')) in allowed, f'{sid}: unsourced numeric token {token!r} in {text!r}')
    require(footer_found, f'{sid}: copyright footer missing')


def values_equal(actual, expected, context):
    require(len(actual) == len(expected), f'{context}: length {len(actual)} != {len(expected)}')
    for i, (a, e) in enumerate(zip(actual, expected)):
        if e is None:
            require(a is None, f'{context}[{i}]: expected blank, got {a}')
        else:
            require(a is not None and abs(D(str(a)) - D(str(e))) <= D('0.000000001'),
                    f'{context}[{i}]: {a} != {e}')


def chart_expected(sid, d, c):
    arr = d['arr_history']['arr']
    if sid in {'BD04', 'IN03', 'IN14'}:
        vs, actual = (arr, 9) if sid == 'IN03' else (arr[6:], 3)
        return [('Actual', [v if j < actual else None for j, v in enumerate(vs)]),
                ('Forecast', [v if j >= actual else None for j, v in enumerate(vs)])]
    if sid == 'BD07':
        vals = [d['plan_vs_actual_q3_fy26'][5]['actual']] + d['cash_runway']['cash_end']
        return [('Actual', [vals[0]] + [None] * (len(vals) - 1)), ('Forecast', [None] + vals[1:])]
    if sid in {'BD03', 'BD09'}:
        if sid == 'BD03':
            b = d['arr_bridge_q3_fy26']
            values = [b[k] for k in ('start', 'new', 'expansion', 'contraction', 'churn', 'end')]
            totals = {0, 5}
        else:
            h = d['headcount']
            values = [sum(h['start']), sum(h['joiners']), -sum(h['leavers']), sum(h['end']), sum(h['plan_q4'])]
            totals = {0, 3, 4}
        base, movement, running = [], [], D(0)
        for index, value in enumerate(values):
            if index in totals:
                base.append(D(0)); movement.append(value); running = value
            else:
                base.append(min(running, running + value)); movement.append(abs(value)); running += value
        return [('Hidden base', base), ('Movement', movement)]
    if sid == 'BD08': return [('Value', [row['min_cash'] for row in d['cases']])]
    if sid == 'BD10': return [(s['stage'], [v]) for s, v in zip(d['pipeline_q4_fy26']['stages'], c['weights'])]
    if sid == 'IN05': return [('Gross margin', d['unit_economics']['gross_margin_trend']['gm'])]
    if sid == 'IN10': return [('Allocation', [row['pct'] for row in d['round']['use_of_funds']])]
    if sid == 'IN07': return [(p['name'], ([p['x']], [p['y']])) for p in d['competitive_2x2']['players']]
    if sid == 'IN08':
        rows = [(p['name'], ([p['growth']], [p['valuation_arr']])) for p in d['comps']]
        rows += [('Parspec', ([c['yoy']], [c['multiple']])),
                 ('Growth median', ([c['growth_median']] * 2, [D(0), D(20)])),
                 ('Valuation median', ([D('.15'), D('.95')], [c['median']] * 2))]
        return rows
    raise AssertionError('unexpected chart page ' + sid)


def category_expected(sid, d):
    if sid in {'BD04', 'IN14'}: return d['arr_history']['quarters'][6:]
    if sid == 'IN03': return d['arr_history']['quarters']
    if sid == 'BD03': return ['Start\nQ3 FY26', 'New logos', 'Expansion', 'Contraction', 'Churn', 'End\nQ3 FY26']
    if sid == 'BD09': return ['Start\nQ3 FY26', 'Joiners', 'Leavers', 'End\nQ3 FY26', 'Plan\nQ4 FY26']
    if sid == 'BD08': return [row['case'] for row in d['cases']]
    if sid == 'BD10': return ['Weighted pipeline']
    if sid == 'IN05': return d['unit_economics']['gross_margin_trend']['quarters']
    if sid == 'IN10': return [row['use'] for row in d['round']['use_of_funds']]
    if sid == 'BD07':
        result = ['Oct 2026\n(Q3 FY26)']
        months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
        for label in d['cash_runway']['months']:
            name, year_text = label.split(); year = 2000 + int(year_text)
            month = months.index(name) + 1
            fy = year if month >= 2 else year - 1
            quarter = ((month - 2) % 12) // 3 + 1
            result.append(f'{name} {year}\n(Q{quarter} FY{str(fy)[-2:]})')
        return result
    return None


def workbook_contract(chart, sid):
    workbook = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob), data_only=True)
    for ref in chart._chartSpace.xpath('.//c:numRef | .//c:strRef'):
        formula = ref.find(qn('c:f')).text
        sheet_name, area = formula.rsplit('!', 1)
        sheet_name = sheet_name.strip("'").replace("''", "'")
        x0, y0, x1, y1 = range_boundaries(area.replace('$', ''))
        values = [workbook[sheet_name].cell(y, x).value for y in range(y0, y1 + 1) for x in range(x0, x1 + 1)]
        cache = ref.find(qn('c:numCache'))
        numeric = cache is not None
        if cache is None: cache = ref.find(qn('c:strCache'))
        require(cache is not None, f'{sid}: cache missing for {formula}')
        cached = {int(pt.get('idx')): pt.find(qn('c:v')).text for pt in cache.findall(qn('c:pt'))}
        require(int(cache.find(qn('c:ptCount')).get('val')) == len(values), f'{sid}: cache/workbook point count')
        for index, value in enumerate(values):
            if value is None:
                require(index not in cached or cached[index] in {'', None}, f'{sid}: nonblank cache for workbook blank')
            elif numeric:
                values_equal([cached.get(index)], [value], f'{sid}: workbook {formula}[{index}]')
            else:
                require(cached.get(index) == str(value), f'{sid}: workbook text/cache mismatch {formula}')


def chart_contract(slide, sid, d, c, allowed):
    charts = [shape for shape in slide.shapes if shape.has_chart]
    require(len(charts) == (1 if sid in CHART_SLIDES else 0), f'{sid}: native chart count {len(charts)}')
    # Chart-looking names must actually reference editable chart objects.
    for shape in slide.shapes:
        if shape.name.startswith('native:') and shape.name.split(':', 1)[1] not in {'fin_table', 'heatmap_table'}:
            require(shape.has_chart, f'{sid}: chart primitive is not a native chart')
    for shape in charts:
        chart = shape.chart
        require(not chart.has_title, f'{sid}: internal chart title')
        require(len(chart._chartSpace.xpath('.//c:valAx/c:majorGridlines')) <= 1, f'{sid}: excessive horizontal gridlines')
        require(not chart._chartSpace.xpath('.//c:valAx/c:minorGridlines'), f'{sid}: minor horizontal gridlines')
        expected = chart_expected(sid, d, c)
        categories = category_expected(sid, d)
        if categories is not None:
            actual_categories = [category.label for category in chart.plots[0].categories]
            require(actual_categories == categories, f'{sid}: chart category / fiscal label mismatch')
            for category in actual_categories: fiscal_contract(category, sid)
        require(len(chart.series) == len(expected), f'{sid}: series count')
        by_name = {series.name: series for series in chart.series}
        require(set(by_name) == {name for name, _ in expected}, f'{sid}: chart series names')
        for name, values in expected:
            series = by_name[name]
            if sid in {'IN07', 'IN08'}:
                x_cache = series._element.find(qn('c:xVal')).find(qn('c:numRef')).find(qn('c:numCache'))
                x_values = [None] * int(x_cache.find(qn('c:ptCount')).get('val'))
                for point in x_cache.findall(qn('c:pt')):
                    x_values[int(point.get('idx'))] = point.find(qn('c:v')).text
                values_equal(x_values, values[0], f'{sid}/{name}/x')
                values_equal(series.values, values[1], f'{sid}/{name}/y')
            else:
                values_equal(series.values, values, f'{sid}/{name}')
        for rp in chart._chartSpace.xpath('.//a:rPr[@sz] | .//a:defRPr[@sz]'):
            require(int(rp.get('sz')) / 100 * 128 / 72 + .02 >= 14, f'{sid}: chart font below 14px')
        require(all(name == FONT for name in chart._chartSpace.xpath('.//a:latin/@typeface')), f'{sid}: nonbrand chart font')
        require(chart._chartSpace.xpath('./c:txPr//a:defRPr[@sz]'), f'{sid}: chart text size not explicit')
        if sid not in {'BD12', 'IN12'}:
            require(not chart._chartSpace.xpath('.//a:rPr[@baseline>0]'), f'{sid}: chart superscript')
        require(chart._chartSpace.xpath('.//c:dLbls'), f'{sid}: native chart data labels missing')
        for text in chart._chartSpace.xpath('.//c:rich//a:t/text()'):
            for token in NUMBER.findall(text):
                require(D(token.replace(',', '')) in allowed, f'{sid}: unsourced chart-label number {token}')
        workbook_contract(chart, sid)


def table_matrix(shape):
    return [[cell.text.replace('\v', '\n') for cell in row.cells] for row in shape.table.rows]


def pnl_expected(d, c, detail=False):
    p = d['pnl']
    rows = [('Revenue', p['revenue']), ('Gross profit', c['gp']),
            ('Gross margin', [g / r for g, r in zip(c['gp'], p['revenue'])]),
            ('Operating expenses', p['opex']), ('EBITDA', c['ebitda'])]
    if detail:
        rows.insert(1, ('Revenue growth, QoQ', [None] + [p['revenue'][i] / p['revenue'][i - 1] - 1 for i in range(1, 4)]))
        rows.insert(2, ('Cost of revenue', c['cor']))
        rows.append(('EBITDA margin', [e / r for e, r in zip(c['ebitda'], p['revenue'])]))
    result = []
    for name, values in rows:
        if name == 'Gross margin': total = sum(c['gp']) / sum(p['revenue'])
        elif name == 'EBITDA margin': total = sum(c['ebitda']) / sum(p['revenue'])
        elif 'growth' in name: total = None
        else: total = sum(values)
        percent = 'margin' in name or 'growth' in name
        result.append([name] + [('—' if value is None else pct(value, 1 if 'growth' in name else 0) if percent else fmt(value, 2 if detail else 1)) for value in [*values, total]])
    return result


def table_contract(slide, sid, d, c):
    tables = [table_matrix(s) for s in slide.shapes if s.has_table]
    require(bool(tables) == (sid in TABLE_SLIDES), f'{sid}: native table presence')
    # Measure the supported native spacing rather than assuming Montserrat's
    # proportional default digits are tabular. Restrict to financial body cells.
    metrics = load_brand().investor['financial_digit_metrics']
    for shape in slide.shapes:
        if not shape.has_table: continue
        for row in list(shape.table.rows)[1:]:
            for cell in list(row.cells)[1:]:
                if not re.fullmatch(r'[+−()0-9.,%× pt—]+', cell.text): continue
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if len(run.text) == 1 and run.text in '0123456789':
                            widths = metrics['bold' if run.font.bold else 'regular']
                            expected = round((max(widths)-widths[int(run.text)]) / metrics['units_per_em'] * run.font.size.pt * 100)
                            require(int(run._r.get_or_add_rPr().get('spc', '0')) == expected,
                                    f'{sid}: digit does not have native tabular advance')

    def match_rows(expected, label):
        require(any(table[1:] == expected for table in tables), f'{sid}: {label} table differs from independent source calculation')
    if sid in {'BD06', 'BD12'}: match_rows(pnl_expected(d, c, sid == 'BD12'), 'P&L')
    if sid == 'BD01a':
        expected = [[str(i + 1), a['item'], a['owner'], str(a['minutes'])] for i, a in enumerate(d['agenda'])]
        expected.append(['', 'Total', '', str(sum(a['minutes'] for a in d['agenda']))])
        match_rows(expected, 'agenda')
    if sid == 'BD05':
        expected = []
        for i, row in enumerate(d['plan_vs_actual_q3_fy26']):
            delta = row['actual'] - row['plan']
            favorable = delta > 0 if i in {0, 1, 2, 5} else delta < 0
            expected.append([row['metric'], pct(row['plan']) if i == 2 else fmt(row['plan']),
                             pct(row['actual']) if i == 2 else fmt(row['actual']),
                             ('+' if delta > 0 else '−') + (fmt(abs(delta) * 100, 0) + ' pt' if i == 2 else fmt(abs(delta))),
                             '▲ Ahead' if favorable else '▼ Behind'])
        match_rows(expected, 'plan versus actual')
    if sid == 'BD09':
        hc = d['headcount']; expected = []
        for i, name in enumerate(hc['functions']):
            expected.append([name, str(hc['start'][i]), '+' + str(hc['joiners'][i]),
                             '(' + str(hc['leavers'][i]) + ')' if hc['leavers'][i] else '—', str(hc['end'][i]), str(hc['plan_q4'][i])])
        expected.append(['Total', str(sum(hc['start'])), '+' + str(sum(hc['joiners'])), '(' + str(sum(hc['leavers'])) + ')', str(sum(hc['end'])), str(sum(hc['plan_q4']))])
        match_rows(expected, 'headcount')
    if sid == 'BD10':
        stages = d['pipeline_q4_fy26']['stages']
        expected = [[s['stage'], fmt(s['amount']), pct(s['win_rate']), fmt(v)] for s, v in zip(stages, c['weights'])]
        expected.append(['Total', fmt(sum(s['amount'] for s in stages)), '—', fmt(c['weighted'])])
        match_rows(expected, 'weighted pipeline')
    if sid == 'BD07':
        burn = d['cash_runway']['net_burn']
        match_rows([['Net burn', fmt(sum(burn)), fmt(sum(burn) / len(burn))]], 'forward burn')
    if sid == 'BD08':
        match_rows([[row['case'], fmt(row['fy27_exit_arr']), fmt(row['fy27_revenue']), row['cash_out']] for row in d['cases']], 'named cases')
    if sid == 'BD13':
        cash = d['cash_runway']; draw = cash['debt_draw']['amount']; values = cash['cash_end']
        expected = [['Cash · Jan 2027 (Q4 FY26)', fmt(values[2]), fmt(values[2])],
                    ['Cash · Feb 2027 (Q1 FY27)', fmt(values[3] - draw), fmt(values[3])],
                    ['Cash · Oct 2027 (Q3 FY27)', fmt(values[-1] - draw), fmt(values[-1])],
                    ['Net burn · next twelve months', fmt(sum(cash['net_burn'])), fmt(sum(cash['net_burn']))]]
        match_rows(expected, 'debt cash comparison')
        match_rows([['Runway', str(cash['runway_months_without_debt']), str(cash['runway_months_with_debt'])]], 'runway units')
    if sid == 'IN04':
        co = d['cohort_ndr']
        expected = [[name] + ['' if v is None else pct(v) for v in values] for name, values in zip(co['cohorts'], co['ndr'])]
        match_rows(expected, 'cohort triangle')
    if sid == 'IN09':
        peers = [(s['name'], s['growth'], s['valuation_arr']) for s in d['comps']] + [('Parspec', c['yoy'], c['multiple'])]
        expected = [[name, pct(growth, 1 if name == 'Parspec' else 0), multiple(v), ('+' if v >= c['median'] else '−') + multiple(abs(v - c['median']))] for name, growth, v in sorted(peers, key=lambda row: row[1], reverse=True)]
        match_rows(expected, 'comparables / unrounded median deltas')
    if sid == 'IN14':
        p, f, arr = d['pnl'], d['fy27_plan'], d['arr_history']['arr']
        rev, gp, nn = p['revenue'] + f['revenue'], c['gp'] + c['gp27'], c['net_new']
        expected = [['Ending ARR'] + [fmt(v) for v in arr[6:]] + [fmt(arr[9]), fmt(arr[-1])],
                    ['Net new ARR'] + [fmt(v) for v in nn] + [fmt(sum(nn[:4])), fmt(sum(nn[4:]))],
                    ['Revenue'] + [fmt(v) for v in rev] + [fmt(sum(rev[:4])), fmt(sum(rev[4:]))],
                    ['Gross profit'] + [fmt(v) for v in gp] + [fmt(sum(gp[:4])), fmt(sum(gp[4:]))],
                    ['Gross margin'] + [pct(g / r) for g, r in zip(gp, rev)] + [pct(sum(gp[:4]) / sum(rev[:4])), pct(sum(gp[4:]) / sum(rev[4:]))]]
        match_rows(expected, 'financial plan')


def staging_contract():
    def git(*args): return subprocess.check_output(['git', *args], cwd=REPO, text=True).splitlines()
    changed = set(git('diff', '--cached', '--name-only')) | set(git('diff', '--name-only', '--diff-filter=A', 'HEAD'))
    for path in changed:
        require(Path(path).suffix.lower() not in {'.pptx', '.ppt', '.png', '.pdf', '.jpg', '.jpeg', '.webp', '.key'}, f'public repo artifact staged/added: {path}')
        require(not re.search(r'(^|/)(?:_refs|templates|renders)(/|$)', path, re.I), f'public repo private asset path: {path}')
    return len(changed)


def mutation_probes(prs, ids, d, c, allowances):
    """Mutate in-memory slide copies and require each guard to reject them."""
    from pptx.slide import Slide
    sid = 'BD04'; source = prs.slides[ids.index(sid)]
    def altered():
        # python-pptx shape proxies keep the copied XML editable; relationships
        # remain valid for read-only chart inspection on the original slide part.
        return Slide(deepcopy(source._element), source.part)
    probes = []
    def rejected(name, slide, checker):
        try: checker(slide)
        except AssertionError: probes.append(name)
        else: raise AssertionError('mutation survived: ' + name)
    slide = altered()
    box = next(shape for shape in slide.shapes if shape.has_text_frame and shape.top / 914400 * 128 < 672 and shape.text.strip())
    box.text_frame.paragraphs[0].runs[0].text = '999999.9'
    rejected('unsourced number', slide, lambda s: text_contract(s, sid, allowances[sid]))
    slide = altered(); box = next(shape for shape in slide.shapes if shape.has_text_frame and shape.top / 914400 * 128 < 672 and shape.text.strip())
    from pptx.util import Pt
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(5)
    rejected('font floor', slide, lambda s: text_contract(s, sid, allowances[sid]))
    slide = altered(); box = next(shape for shape in slide.shapes if shape.has_text_frame and shape.top / 914400 * 128 < 672 and shape.text.strip())
    box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set('baseline', '30000')
    rejected('superscript', slide, lambda s: text_contract(s, sid, allowances[sid]))
    slide = altered(); box = next(shape for shape in slide.shapes if shape.has_text_frame and shape.top / 914400 * 128 < 672 and shape.text.strip())
    box.text_frame.paragraphs[0].runs[0].text = 'Q3 performance'
    rejected('bare fiscal quarter', slide, lambda s: text_contract(s, sid, allowances[sid]))
    slide = altered(); chart = next(shape for shape in slide.shapes if shape.has_chart)
    chart._element.getparent().remove(chart._element)
    rejected('missing native chart', slide, lambda s: chart_contract(s, sid, d, c, allowances[sid]))
    # A source-plausible number still fails when it occupies the wrong table
    # cell; this guards against a permissive numeric-token whitelist.
    table_sid = 'BD05'
    table_slide = Slide(deepcopy(prs.slides[ids.index(table_sid)]._element), prs.slides[ids.index(table_sid)].part)
    table = next(shape.table for shape in table_slide.shapes if shape.has_table)
    table.cell(1, 2).text = fmt(d['plan_vs_actual_q3_fy26'][0]['plan'])
    rejected('wrong source-plausible financial cell', table_slide, lambda s: table_contract(s, table_sid, d, c))
    # A copied slide retains chart-part relationships. Restore this in-memory
    # chart cache immediately after the negative probe; no file is written.
    slide = altered()
    chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
    value = chart._chartSpace.xpath('.//c:ser[1]/c:val/c:numRef/c:numCache/c:pt[1]/c:v')[0]
    original = value.text
    try:
        value.text = '999999.9'
        rejected('corrupt native chart cache', slide, lambda s: chart_contract(s, sid, d, c, allowances[sid]))
    finally:
        value.text = original
    return probes


def verify(path, ids, d, c, allowances, sha):
    prs = Presentation(path)
    require(len(prs.slides) == len(ids), f'{path.name}: slide count')
    charts = tables = 0
    for slide, sid in zip(prs.slides, ids):
        notes = slide.notes_slide.notes_text_frame.text
        require(f'Slide: {sid}' in notes, f'{path.name}: ID/order mismatch {sid}')
        require(sha in notes, f'{sid}: fixture hash provenance missing')
        text_contract(slide, sid, allowances[sid])
        chart_contract(slide, sid, d, c, allowances[sid])
        table_contract(slide, sid, d, c)
        charts += sum(shape.has_chart for shape in slide.shapes)
        tables += sum(shape.has_table for shape in slide.shapes)
    return prs, dict(path=str(path), slides=len(ids), native_charts=charts, native_tables=tables,
                     sha256=hashlib.sha256(path.read_bytes()).hexdigest())


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--data', type=Path, required=True)
    ap.add_argument('--board', type=Path, default=TEMPLATE_DIR / 'parspec_board_gallery.pptx')
    ap.add_argument('--investor', type=Path, default=TEMPLATE_DIR / 'parspec_investor_gallery.pptx')
    ap.add_argument('--evidence', type=Path, help='Optional private JSON verification receipt')
    ap.add_argument('--self-test', action='store_true', help='Mutation probes also run by default')
    args = ap.parse_args()
    raw = args.data.read_bytes(); d = json.loads(raw, parse_float=D)
    require('FICTIONAL' in d.get('_note', '') and 'v3' in d.get('_note', ''), 'explicit fictional v3 fixture required')
    require(DEFAULT_PROFILE == 'master', 'master must remain the default')
    sha = hashlib.sha256(raw).hexdigest(); c = independent(d)
    allowances = numeric_allowances(d, c)
    board, b = verify(args.board, BOARD_IDS, d, c, allowances, sha)
    _, i = verify(args.investor, INVESTOR_IDS, d, c, allowances, sha)
    mutations = mutation_probes(board, BOARD_IDS, d, c, allowances)
    staged = staging_contract()
    receipt = dict(result='PASS', fixture_sha256=sha, decks=[b, i], mutation_probes=mutations,
                   staged_or_added_paths_checked=staged,
                   limitations='Structural/data verification; visual review, PowerPoint fidelity and PM/brand acceptance are separate.')
    if args.evidence:
        args.evidence.parent.mkdir(parents=True, exist_ok=True)
        args.evidence.write_text(json.dumps(receipt, indent=2) + '\n')
    print(f'PASS: {b["slides"] + i["slides"]} slides; {b["native_charts"] + i["native_charts"]} native charts with source-exact caches and coherent workbooks; financial tables; per-slide numeric provenance; text/fiscal/source guards; {len(mutations)} rejected mutations; staging guard')


if __name__ == '__main__':
    main()
