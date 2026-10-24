"""Synthetic regression checks for the opt-in proposed sales register.

Run from any directory. --gallery adds editable-object, copy and text-floor
checks against a locally built deck. This does not certify visual fidelity;
inspect the LibreOffice renders separately. No company template is distributed.
"""
from __future__ import annotations

import argparse
import math
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from types import SimpleNamespace

from pptx import Presentation
import yaml

SKILL = Path(__file__).resolve().parents[1]
REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(SKILL))
from assets.brand import load, _design_model_path, _resolve, _hex_to_rgb
from assets.house import DEFAULT_PROFILE, PROFILES
from assets.motif import Motif


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def contrast_ratio(foreground, background):
    """WCAG sRGB relative-luminance contrast, using the actual role colors."""
    def luminance(rgb):
        channels = [c / 255 for c in rgb]
        linear = [c / 12.92 if c <= .04045 else ((c + .055) / 1.055) ** 2.4 for c in channels]
        return sum(c * weight for c, weight in zip(linear, (.2126, .7152, .0722)))
    a, b = sorted((luminance(foreground), luminance(background)))
    return (b + .05) / (a + .05)


def model_contract(base_ref):
    model = _design_model_path()
    current = yaml.safe_load(model.read_text())
    rel = str(model.relative_to(REPO))
    original = yaml.safe_load(subprocess.check_output(
        ["git", "show", f"{base_ref}:{rel}"], cwd=REPO, text=True))
    permitted = {"meta.version", "meta.status", "meta.updated", "meta.supersedes", "meta.notes"}

    def walk(before, after, path=""):
        if isinstance(before, dict):
            require(isinstance(after, dict), f"removed mapping: {path}")
            for key, value in before.items():
                child = f"{path}.{key}" if path else str(key)
                require(key in after, f"removed existing key: {child}")
                walk(value, after[key], child)
        elif path not in permitted:
            require(before == after, f"changed existing value: {path}")

    walk(original, current)
    # Meta history can be prepended, but earlier explanation must survive.
    require(original["meta"]["notes"].strip() in current["meta"]["notes"], "existing meta history removed")
    require(tuple(int(v) for v in current["meta"]["version"].split(".")) >= (2, 3, 0), "model predates the sales register")
    require(current["meta"]["status"] == "proposed", "proposal marked ratified")
    require(current["meta"]["brand_version"] == original["meta"]["brand_version"], "brand version changed")
    require(DEFAULT_PROFILE == "master", "sales silently became default")
    require(PROFILES["sales"].title_bold and not PROFILES["sales"].title_caps, "sales title register")
    b = load()
    require(b.sales["geometry"]["px_per_in"] == 128, "sales conversion basis")
    require(b.sales["geometry"]["content_top_px"] == 184, "sales content band")
    for key, raw in current["sales_register"]["colors"].items():
        require(b.role("sales_" + key) == _hex_to_rgb(_resolve(raw, current)), f"unresolved sales color: {key}")
    for name in ("eyebrow", "title"):
        spec = b.sales_type(name)
        require(spec["weight"] == 700 and spec["letter_spacing"] == 0, f"tracked/nonbold {name}")
    require(b.sales_type("title")["size_px"] == 42, "sales title scale")
    print("PASS model: additive existing keys, unchanged brand/default, proposed status, resolved roles")
    return b


def primitive_contract(b):
    prs = Presentation()
    prs.slide_width, prs.slide_height = 9144000, 5143500
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    motif = Motif(SimpleNamespace(brand=b))
    title = motif.text(slide, 56, 70, 1168, 57, "Synthetic title", style="title")
    run = title.text_frame.paragraphs[0].runs[0]
    require(math.isclose(run.font.size.pt, 42 * 72 / 128, abs_tol=.01), "42px did not become 23.625pt")
    require(run.font.name == b.font, "title family does not come from brand")
    label = motif.text(slide, 56, 130, 400, 24, "Minimum label", style="label")
    require(label.text_frame.paragraphs[0].runs[0].font.size.pt * 128 / 72 >= 14,
            "centipoint conversion lowered the 14px floor")
    light = motif.text(slide, 56, 150, 400, 60, "Agenda", style="agenda_display")
    require(light.text_frame.paragraphs[0].runs[0].font.name == b.sales["native_font_faces"]["light"], "light token compiled as regular")
    bold = motif.text(slide, 56, 220, 100, 50, "FF", style="small_numeral", bold=True)
    require(bold.text_frame.paragraphs[0].runs[0].font.name == b.font, "bold override inherited light face")
    items = [{"label": str(i), "value": n} for i, n in enumerate([30, 18, 7, 6, 4], 1)]
    n_before = len(slide.shapes)
    try:
        motif.donut_callouts(slide, items, total=64)
    except ValueError:
        pass
    else:
        raise AssertionError("mismatched donut sum accepted")
    require(len(slide.shapes) == n_before, "rejected donut partially mutated slide")
    try:
        motif.donut_callouts(slide, [{"label": "empty", "value": 0}])
    except ValueError:
        pass
    else:
        raise AssertionError("zero donut accepted")
    try:
        motif.color("missing_role")
    except KeyError:
        pass
    else:
        raise AssertionError("unknown role silently substituted")
    motif.donut_callouts(slide, items, total=65)
    require(sum(s._element.xpath('.//a:prstGeom[@prst="blockArc"]') != [] for s in slide.shapes) == 5,
            "donut is not five editable native segments")
    with tempfile.TemporaryDirectory(prefix="sales-smoke-") as tmp:
        path = Path(tmp) / "synthetic.pptx"
        prs.save(path)
        reopened = Presentation(path)
        require(len(reopened.slides) == 1, "synthetic deck did not round-trip")
        require(sum(bool(s._element.xpath('.//a:prstGeom[@prst="blockArc"]')) for s in reopened.slides[0].shapes) == 5,
                "native segments lost on round-trip")
    print("PASS primitives: px/pt + font mapping, five native segments, rejected bad sums/zero/unknown role")


def gallery_contract(path, b):
    prs = Presentation(path.expanduser())
    require(len(prs.slides) == 10, "gallery must contain exactly ten slides")
    require(prs.slide_width == 9144000 and prs.slide_height == 5143500, "gallery canvas must be 10x5.625in")
    require(len(prs.slide_masters) == 1, "gallery retained extra masters")
    master = prs.slide_masters[0]
    logos = [s for s in master.shapes if s.shape_type == 13]
    require(len(logos) == 1 and logos[0].image.size[0] >= 1000, "footer needs real full-resolution logo")
    require(abs(logos[0].left / 914400 * 128 - 56) < .1, "footer logo inset")
    copyright_shapes = [s for s in master.shapes if s.has_text_frame and "©2026" in s.text]
    require(len(copyright_shapes) == 1, "copyright footer missing/duplicated")
    cp = copyright_shapes[0]
    require(abs((cp.left + cp.width / 2) / 914400 * 128 - 640) < .1, "copyright not centered")
    texts = ["\n".join(s.text for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
    for i, slide in enumerate(prs.slides):
        require(slide.background.fill.fore_color.rgb == b.role("sales_ground_master"), f"S{i+1:02d} ground")
        if i:
            require(any(s.has_text_frame and s.text == str(i+1) and s.top / 914400 * 128 >= 672
                        for s in slide.shapes), f"S{i+1:02d} page footer")
    require("45 minutes" in texts[1] and texts[1].count("10 min") == 4 and "5 min" in texts[1], "agenda canonical timings")
    require(texts[2].count("02") == 1, "divider chapter duplicated")
    require(all(x in texts[3] for x in ("20XX", "00", "000+", "$00B", "Sample figures for layout review.")), "company sample copy")
    require("+18%" not in texts[4] and "+1.2" not in texts[4] and "2×" not in texts[4], "rejected proof statistics reintroduced")
    require(sum(bool(s._element.xpath('.//a:prstGeom[@prst="blockArc"]')) for s in prs.slides[5].shapes) == 5,
            "S06 native segment editability")
    require("74%" in texts[5] and "Illustrative sample data." in texts[5], "donut takeaway/disclosure")
    require("Contract to Every Branch Live" in texts[6] and "Pilot (45 days)" in texts[6], "pilot scope copy")
    require(not any(x in texts[6] for x in ("December", "November", "Go-Live in 45 Days")), "unsupported implementation date")
    require(all(x in texts[7] for x in ("Guided, 2 weeks", "All modules", "Dedicated CSM, 45-day pilot", "$24K", "$60K", "$150K", "Sample pricing for layout review.")), "commercial sample copy")
    require("Q3 vs Q2 · illustrative sample data." in texts[8] and "for the first time" not in texts[8].lower(), "QBR comparison basis")
    require(all(x in texts[9] for x in ("Northwind · TBD", "Both · TBD", "Parspec · TBD", "[Executive Name]", "[Account Executive]", "Template-only")), "closer template owners/placeholders")
    deltas = [s for s in prs.slides[8].shapes if s.has_text_frame and s.text.startswith(("▲", "▼"))]
    require(len(deltas) == 6, "QBR needs six delta fields")
    for delta in deltas:
        for para in delta.text_frame.paragraphs:
            for run in para.runs:
                require(run.font.size.pt * 128 / 72 >= 14, "QBR delta below 14px")
                require(run.font.color.rgb == b.role("sales_status_good"), "QBR delta status role")
                require(contrast_ratio(run.font.color.rgb, b.role("sales_card")) >= 4.5, "QBR delta contrast below 4.5:1")
    note = next(s for s in prs.slides[9].shapes if s.has_text_frame and s.text.startswith("Template-only"))
    arcs = [s for s in prs.slides[9].shapes if s.width / 914400 * 128 > 200 and
            s._element.xpath('.//a:prstGeom[@prst="ellipse"]')]
    require(arcs and note.top + note.height <= min(s.top for s in arcs), "template note enters quarter-arc bounding box")
    floor = b.sales["minimum_meaningful_text_px"]
    small_runs = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text.strip():
                continue
            px_top = shape.top / 914400 * 128
            exception = (px_top >= 672 or re.match(r"^\d+\.", shape.text) or
                         shape.text == "Q3 vs Q2 · illustrative sample data." or
                         (shape.text.isdigit() and len(shape.text) == 1))
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip() and run.font.size is not None:
                        size_px = run.font.size.pt * 128 / 72
                        if size_px < floor - .1 and not exception:
                            small_runs.append(f"S{i+1:02d}:{run.text[:36]} ({size_px:g}px)")
    require(not small_runs, "meaningful text below floor: " + "; ".join(small_runs))
    print("PASS gallery: 10 slides, local real-logo footer, editable donut, canonical copy, 14px content floor")


def render_contract(gallery, pdf, b):
    """Check actual rendered lines; text-presence alone cannot detect wrapping."""
    import fitz
    prs = Presentation(gallery.expanduser())
    doc = fitz.open(pdf.expanduser())
    require(len(doc) == 10, "render must contain ten pages")
    def norm(text):
        return re.sub(r"\s+", "", text)
    def lines(index):
        return [dict(text="".join(s["text"] for s in line["spans"]),
                     bbox=fitz.Rect(line["bbox"]), spans=line["spans"])
                for block in doc[index].get_text("dict")["blocks"]
                for line in block.get("lines", [])]
    def exact(index, text):
        return [line for line in lines(index) if norm(line["text"]) == norm(text)]
    label_rows = []
    for label in ("FOUNDED", "TEAM", "DISTRIBUTORS & REPS", "QUOTED VALUE"):
        found = exact(3, label)
        require(len(found) == 1, f"S04 label wraps/missing: {label}")
        label_rows.append(found[0])
    markers = exact(3, "1")
    require(len(markers) == 4, "S04 needs four source markers")
    require(not any((label["bbox"] & marker["bbox"]).get_area() > 0
                    for label in label_rows for marker in markers), "S04 label/marker collision")
    body_cells = [s for s in prs.slides[7].shapes if s.has_text_frame and s.text.strip()
                  and 244 <= s.top / 914400 * 128 < 584]
    cell_boxes = [fitz.Rect(s.left / 12700, s.top / 12700,
                           (s.left + s.width) / 12700, (s.top + s.height) / 12700)
                  for s in prs.slides[7].shapes if s._element.xpath('.//a:prstGeom[@prst="rect"]')
                  and not (s.has_text_frame and s.text.strip())]
    require(len(body_cells) == 20, "S08 body-cell coverage")
    for cell in body_cells:
        rect = fitz.Rect(cell.left / 12700, cell.top / 12700,
                         (cell.left + cell.width) / 12700, (cell.top + cell.height) / 12700)
        found = [line for line in exact(7, cell.text) if rect.contains(
            fitz.Point((line["bbox"].x0 + line["bbox"].x1) / 2,
                       (line["bbox"].y0 + line["bbox"].y1) / 2))]
        require(len(found) == 1, f"S08 cell wraps/missing: {cell.text}")
        # Use the actual cell panel: right-aligned glyph bearings can extend
        # slightly beyond their text-frame advance while staying inside padding.
        panels = [panel for panel in cell_boxes if panel.contains(rect)]
        require(panels and min(panels, key=lambda panel: panel.get_area()).contains(found[0]["bbox"]),
                f"S08 text overflows its cell: {cell.text}")
        require(all(s["size"] * 128 / 72 >= 17 for s in found[0]["spans"]), "S08 body shrunk below 17px")
    delta_rows = [line for line in lines(8) if line["text"].startswith(("▲", "▼"))]
    require(len(delta_rows) == 6, "rendered QBR delta coverage")
    min_size = min(s["size"] * 128 / 72 for row in delta_rows for s in row["spans"])
    contrasts = [contrast_ratio(((s["color"] >> 16) & 255, (s["color"] >> 8) & 255,
                                 s["color"] & 255), b.role("sales_card"))
                 for row in delta_rows for s in row["spans"]]
    require(min_size >= 14 and min(contrasts) >= 4.5, "rendered QBR delta size/contrast")
    note = exact(9, "Template-only contact and headshot placeholders")
    arcs = [s for s in prs.slides[9].shapes if s.width / 914400 * 128 > 200 and
            s._element.xpath('.//a:prstGeom[@prst="ellipse"]')]
    require(len(note) == 1 and note[0]["bbox"].y1 <= min(s.top / 12700 for s in arcs),
            "rendered template note enters quarter-arc bounding box")
    print(f"PASS render: S04 four single-line labels; S08 twenty single-line cells >=17px; "
          f"S10 note above arcs; S09 deltas >= {min_size:.2f}px, contrast {min(contrasts):.2f}:1")


def staging_contract():
    staged = subprocess.check_output(["git", "diff", "--cached", "--name-only"], cwd=REPO, text=True).splitlines()
    forbidden = {".pptx", ".ppt", ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".key"}
    require(not [p for p in staged if Path(p).suffix.lower() in forbidden], "proprietary binary staged")
    print("PASS staging: no decks, templates or renders staged")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--base-ref", default="HEAD", help="Pre-upgrade git revision for additive-key regression")
    ap.add_argument("--gallery", type=Path)
    ap.add_argument("--pdf", type=Path, help="Rendered gallery PDF; enables line-fit/contrast regression checks")
    args = ap.parse_args()
    b = model_contract(args.base_ref)
    primitive_contract(b)
    if args.gallery:
        gallery_contract(args.gallery, b)
    if args.pdf:
        if not args.gallery:
            ap.error("--pdf requires --gallery")
        render_contract(args.gallery, args.pdf, b)
    staging_contract()


if __name__ == "__main__":
    main()
