"""
parspec-pptx — diagram + composition primitives.

Canvas-agnostic: every coordinate is a float in inches, so the same
primitives work on the 13.333" brand_artifacts canvas and the 10" town hall
canvas. Colors come from a `Palette` resolved off design-model.yaml roles.

Primitives
    text()             multi-paragraph, multi-run text box (or fills a placeholder)
    shape()            autoshape with optional fill/line alpha, rotation, adjustments
    card()             dark card with optional left accent bar
    kpi_tile()         label / big value / delta line
    chip_row()         N equal cards: bold headline + muted detail
    funnel()           stacked trapezoids, alpha-stepped, value inside + labels right
    process()          pentagon + chevrons with a "we are here" marker
    question_cards()   stacked quote cards
    evidence_rows()    question card -> arrow -> answer card
    section_strip()    orange rule + numbered labels (divider slides)
    arcs()             triple-arc signature motif, bleeding off the top-right
"""

from __future__ import annotations

from dataclasses import dataclass

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .brand import Brand, load as load_brand

FONT = "Montserrat"
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


@dataclass(frozen=True)
class Palette:
    """Dark-surface composition palette. Every value is a design-model role.

    Values observed in the Q2-26 town hall deck are noted for traceability;
    each is within a shade of the role it resolves to.
    """
    bg: RGBColor           # surface_dark          (town hall #0F0F0F)
    card: RGBColor         # surface_dark_alt      (town hall #1A1A1D)
    card_line: RGBColor    # surface_dark_panel    (town hall #2A2A2E)
    accent: RGBColor       # cta                   (town hall #FFA72B)
    accent_deep: RGBColor  # accent_type_on_light  (town hall #8C5E22 funnel stroke)
    text: RGBColor         # text_dark_primary
    muted: RGBColor        # text_dark_tertiary    (town hall #9CA3AF / #9E9E9E)
    good: RGBColor         # status_success_on_dark (town hall #4ADE80)

    @classmethod
    def from_brand(cls, b: Brand | None = None) -> "Palette":
        b = b or load_brand()
        return cls(
            bg=b.surface_dark, card=b.surface_dark_alt, card_line=b.surface_dark_panel,
            accent=b.cta, accent_deep=b.accent_type_on_light, text=b.text_dark_primary,
            muted=b.text_dark_tertiary, good=b.status_success_on_dark,
        )


# ── base primitives ──────────────────────────────────────────────────
def _alpha(parent, pct: float):
    clr = parent.find(qn("a:srgbClr"))
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(pct * 1000)))


def shape(slide, kind, x, y, w, h, *, fill=None, line=None, lw=0.75,
          fill_alpha=None, line_alpha=None, rot=None, adj=None):
    """Autoshape. fill/line None = no fill/line. Alphas are 0–100 (% opaque)."""
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if fill_alpha is not None:
            _alpha(sh._element.spPr.find(qn("a:solidFill")), fill_alpha)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if line_alpha is not None:
            _alpha(sh._element.spPr.find(qn("a:ln")).find(qn("a:solidFill")), line_alpha)
    if rot is not None:
        sh.rotation = rot
    for i, v in enumerate(adj or ()):
        sh.adjustments[i] = v
    return sh


def _bullet(p, color: RGBColor, indent=0.16):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent)))
    pPr.set("indent", str(-Inches(indent)))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def text(slide, x, y, w, h, paras, *, size=11, color=None, bold=False,
         italic=False, align="l", anchor="t", space_after=0, bullet=None,
         line_spacing=None, font=FONT, target=None):
    """Text box with zero insets.

    paras: "one\\ntwo" | ["para", [(run_text, {size,color,bold,italic,font}), ...]]
    bullet: RGBColor for an orange/grey bullet on every paragraph.
    target: an existing shape/placeholder to fill instead of a new text box.
    """
    color = color or RGBColor(0xFF, 0xFF, 0xFF)
    tb = target or slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    if isinstance(paras, str):
        paras = paras.split("\n")
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if bullet is not None:
            _bullet(p, bullet)
        if space_after and i < len(paras) - 1:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def placeholder(t: str, pal: Palette):
    """Run style for content the presenter still has to supply: '[Round size]'."""
    return (t, {"color": pal.muted, "italic": True})


def card(slide, x, y, w, h, pal: Palette, *, accent="accent"):
    """Dark card. accent: 'accent' (orange bar), 'muted' (grey bar) or None."""
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=pal.card, line=pal.card_line)
    if accent:
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.06, h, fill=getattr(pal, accent))


def label(slide, x, y, w, s, pal: Palette, *, color=None, size=8):
    """Small caps section label ('WHAT DOESN'T CHANGE')."""
    return text(slide, x, y, w, 0.2, s.upper(), size=size, bold=True, color=color or pal.accent)


# ── composites ───────────────────────────────────────────────────────
def kpi_tile(slide, x, y, w, h, lab, value, delta, pal: Palette, *, value_size=30):
    """delta: str, or runs [(text, {...})]; use {'color': pal.good} for the good part."""
    card(slide, x, y, w, h, pal)
    label(slide, x + 0.26, y + 0.2, w - 0.4, lab, pal, color=pal.muted)
    text(slide, x + 0.26, y + 0.46, w - 0.4, 0.55, value, size=value_size, bold=True,
         color=pal.accent)
    runs = [(delta, {})] if isinstance(delta, str) else delta
    text(slide, x + 0.26, y + h - 0.4, w - 0.4, 0.22, [runs], size=9, color=pal.muted)


def chip_row(slide, x, y, w, h, items, pal: Palette, *, gap=0.15, head_size=12.5):
    """items: [(headline, detail)] as equal-width accent-less cards."""
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, (head, detail) in enumerate(items):
        cx = x + i * (cw + gap)
        card(slide, cx, y, cw, h, pal, accent=None)
        text(slide, cx + 0.18, y + 0.18, cw - 0.34, 0.5, head, size=head_size, bold=True,
             color=pal.text)
        text(slide, cx + 0.18, y + 0.76, cw - 0.34, h - 0.9, detail, size=9, color=pal.muted)


def funnel(slide, levels, pal: Palette, *, cx=2.55, top=1.42, h=0.8, gap=0.08,
           widths=(4.1, 3.35, 2.6, 1.85), label_x=5.0, label_w=4.5):
    """Narrowing funnel. levels: [(value, label, detail)], widest first.

    Upper levels are accent at stepped alpha (town hall 10/18/32%); the last
    level is solid accent with dark type — the stage the slide is about.
    """
    alphas = [10, 18, 32, 45, 55][: len(levels) - 1] + [None]
    for i, (value, lab, detail) in enumerate(levels):
        w, y, last = widths[i], top + i * (h + gap), i == len(levels) - 1
        shape(slide, MSO_SHAPE.TRAPEZOID, cx - w / 2, y, w, h, fill=pal.accent,
              fill_alpha=alphas[i], line=pal.accent_deep, lw=1.0, rot=180, adj=[0.2])
        text(slide, cx - 0.6, y, 1.2, h, value, size=24, bold=True,
             color=pal.bg if last else pal.text, align="c", anchor="m")
        text(slide, label_x, y + 0.1, label_w, 0.3, lab, size=13, bold=True,
             color=pal.accent if last else pal.text)
        text(slide, label_x, y + 0.42, label_w, 0.3, detail, size=9, color=pal.muted)


def process(slide, stages, current, pal: Palette, *, x=0.44, y=1.5, step=1.81,
            w=1.95, h=0.72, marker="WE ARE HERE"):
    """Pentagon + chevrons. stages: [(name, detail)]. current: index or None.

    Done stages: accent at 35% with a deep stroke. Current: solid accent,
    dark type. Upcoming: card fill, muted type. Numbers are shown because
    a process is a real sequence.
    """
    for i, (name, detail) in enumerate(stages):
        sx = x + i * step
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        if current is not None and i < current:
            style = dict(fill=pal.accent, fill_alpha=35, line=pal.accent_deep)
            tc = pal.text
        elif i == current:
            style, tc = dict(fill=pal.accent), pal.bg
        else:
            style, tc = dict(fill=pal.card, line=pal.card_line), pal.muted
        shape(slide, kind, sx, y, w, h, **style)
        pad = 0.18 if i == 0 else 0.42
        text(slide, sx + pad, y, w - pad - 0.3, h,
             [[(f"{i + 1:02d}  ", {"size": 9}), (name, {})]],
             size=12, bold=True, color=tc, anchor="m")
        done_or_now = current is not None and i <= current
        text(slide, sx + 0.12, y + h + 0.18, w - 0.35, 0.7, detail, size=9,
             color=pal.text if done_or_now else pal.muted, line_spacing=1.1)
    if current is not None and marker:
        mx = x + current * step + w / 2 - 0.1
        shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, mx - 0.09, y + h + 0.98, 0.18, 0.14,
              fill=pal.accent)
        text(slide, mx - 0.8, y + h + 1.16, 1.6, 0.2, marker, size=8, bold=True,
             color=pal.accent, align="c")


def question_cards(slide, x, y, w, items, pal: Palette, *, h=1.1, gap=0.1):
    """items: [(question, context)] stacked, grey bar + accent quote mark."""
    for i, (q, ctx) in enumerate(items):
        cy = y + i * (h + gap)
        card(slide, x, cy, w, h, pal, accent="muted")
        text(slide, x + 0.24, cy + 0.08, 0.5, 0.5, "“", size=34, bold=True, color=pal.accent)
        text(slide, x + 0.72, cy + 0.2, w - 0.95, 0.35, q, size=13, bold=True, color=pal.text)
        text(slide, x + 0.72, cy + 0.6, w - 0.95, 0.35, ctx, size=9, color=pal.muted)


def evidence_rows(slide, x, y, w, rows, pal: Palette, *, q_w=2.45, h=1.05, gap=0.14):
    """rows: [(question, answer_headline, answer_detail)] as Q card -> arrow -> A card."""
    a_w = w - q_w - 0.55
    for i, (q, head, detail) in enumerate(rows):
        cy = y + i * (h + gap)
        card(slide, x, cy, q_w, h, pal, accent="muted")
        text(slide, x + 0.24, cy, q_w - 0.4, h, q, size=13, bold=True, color=pal.text,
             anchor="m")
        shape(slide, MSO_SHAPE.RIGHT_ARROW, x + q_w + 0.13, cy + h / 2 - 0.11, 0.3, 0.22,
              fill=pal.accent)
        ax = x + q_w + 0.55
        card(slide, ax, cy, a_w, h, pal)
        text(slide, ax + 0.26, cy + 0.2, a_w - 0.45, 0.3, head, size=12.5, bold=True,
             color=pal.accent)
        text(slide, ax + 0.26, cy + 0.54, a_w - 0.45, 0.45, detail, size=9.5, color=pal.text,
             line_spacing=1.1)


def section_strip(slide, items, pal: Palette, *, active=None, x=0.18, y=4.43, w=9.76):
    """Numbered agenda strip under an accent rule. active dims every other item."""
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.02, fill=pal.accent)
    colw = w / len(items)
    for i, lab in enumerate(items):
        cx = x + i * colw
        dim = active is not None and i != active
        text(slide, cx, y + 0.13, colw - 0.1, 0.2, f"{i + 1:02d}", size=9, bold=True,
             color=pal.muted if dim else pal.accent)
        text(slide, cx, y + 0.35, colw - 0.1, 0.3, lab, size=10.5,
             color=pal.muted if dim else pal.text)


def arcs(slide, pal: Palette, *, x=7.795, y=-1.848, d=4.044):
    """Triple-arc motif (outer stroke, faint inner stroke, 42% disc), top-right bleed.

    Defaults are the town hall geometry on a 10" canvas; shift x by +3.333
    for the 13.333" canvas.
    """
    k = d / 4.044
    shape(slide, MSO_SHAPE.OVAL, x, y, d, d, line=pal.accent, lw=0.53)
    shape(slide, MSO_SHAPE.OVAL, x + 0.367 * k, y + 0.368 * k, 3.309 * k, 3.309 * k,
          line=pal.accent, lw=0.27, line_alpha=42)
    shape(slide, MSO_SHAPE.OVAL, x + 0.735 * k, y + 0.145 * k, 3.309 * k, 3.309 * k,
          fill=pal.accent, fill_alpha=42)
