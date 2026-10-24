"""Opt-in investor chrome. Geometry and type resolve from design-model.yaml."""
from __future__ import annotations

import math
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from . import shapes as sh


class Investor:
    def __init__(self, deck):
        self.deck, self.brand = deck, deck.brand
        self.spec = self.brand.investor
        self.geo, self.types = self.spec['geometry'], self.spec['type']
        self.ppi = self.geo['px_per_in']

    def color(self, role):
        return self.brand.role('investor_' + role)

    def text(self, slide, x, y, w, h, text, *, kind='body', size=None,
             color='text', bold=None, align='l', anchor='t', leading=None):
        spec = self.types[kind]
        size = spec['size_px'] if size is None else size
        points = math.ceil(size * 72 / self.ppi * 100) / 100
        box = sh.text(slide, *[v/self.ppi for v in (x,y,w,h)], text,
                      size=points, color=self.color(color), font=spec['family'],
                      bold=spec['weight'] >= 600 if bold is None else bold,
                      align=align, anchor=anchor)
        for p in box.text_frame.paragraphs:
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = Pt(points * (leading or spec['line_height']))
        return box

    def rect(self, slide, x,y,w,h, *, fill='card', line=None, lw=1):
        return sh.shape(slide, MSO_SHAPE.RECTANGLE,
                        *[v/self.ppi for v in (x,y,w,h)],
                        fill=self.color(fill) if fill else None,
                        line=self.color(line) if line else None,
                        lw=lw*72/self.ppi)

    def rule(self, slide, x,y,w, *, color='line', h=1):
        return self.rect(slide,x,y,w,h,fill=color)

    def header(self, slide, eyebrow, title, subtitle=''):
        """Explicit line breaks reflow the content band, never shrink the title."""
        self.text(slide,56,40,1168,24,eyebrow.upper(),kind='eyebrow',color='accent')
        lines = title.count('\n') + 1
        self.text(slide,56,70,1168,52*lines,title,kind='title')
        if subtitle:
            self.text(slide,56,128+48*(lines-1),1168,35,subtitle,
                      kind='subtitle',color='muted')
        return 184+48*(lines-1)

    def footer(self, slide, number):
        self.deck.logo(slide,x=56/self.ppi,y=688/self.ppi,w=96/self.ppi)
        self.text(slide,190,688,900,20,
                  '©2026 Parspec Inc. All rights reserved. Proprietary and confidential.',
                  kind='footer',color='muted',align='c')
        self.text(slide,1152,686,72,24,str(number),size=14,
                  color='muted',align='r')

    def tile(self, slide,x,y,w,h,label,value,context='', *, forecast=False,size=None):
        self.rect(slide,x,y,w,h)
        self.rule(slide,x,y,w,color='accent',h=4)
        self.text(slide,x+20,y+22,w-40,38,label.upper(),kind='label',color='muted')
        self.text(slide,x+12,y+65,w-24,76,value,kind='stat',size=size,
                  color='accent' if forecast else 'text')
        if context:
            context_y = max(y+65+(size or self.types['stat']['size_px'])*1.05+10, y+h-46)
            self.text(slide,x+20,context_y,w-40,38,context,size=14,color='muted')

    def callout(self,slide,text,*,x=56,y=584,w=1168,h=56,ask=False):
        self.rect(slide,x,y,w,h,fill='forecast_fill' if ask else 'band')
        self.rect(slide,x,y,4,h,fill='accent')
        self.text(slide,x+24,y+10,w-48,h-18,text,bold=True,anchor='m')

    def cover(self,slide, *, board=False,partner=''):
        self.deck.logo(slide,x=56/self.ppi,y=48/self.ppi,w=176/self.ppi)
        if partner:
            self.rect(slide,252,48,1,30,fill='line')
            self.text(slide,274,50,750,30,partner,bold=True)
        if board:
            self.text(slide,56,232,1000,30,'QUARTERLY BOARD MEETING',kind='eyebrow',color='accent')
            self.text(slide,56,276,1000,164,'Q3 FY26\nBoard Meeting',kind='cover')
            self.rule(slide,56,454,192,color='accent',h=3)
            self.text(slide,56,484,940,32,'Parspec, Inc. standalone · Board of Directors',size=22)
            self.text(slide,56,526,940,32,'Illustrative board template · Confidential',color='muted')
        else:
            self.text(slide,56,214,900,36,'Parspec, Inc. standalone',size=26,bold=True)
            self.text(slide,56,266,1168,132,'Series B',kind='investor_cover')
            self.text(slide,56,412,1168,42,'AI-native operating system for electrical distribution',size=28,color='accent')
            self.rule(slide,56,470,192,color='accent',h=3)
            self.text(slide,56,502,900,32,'Illustrative investor template · Confidential',color='muted')
