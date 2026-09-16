"""
parspec-pptx — house-template register.

Builds decks on a company deck used as a template (its master: PARSPEC logo,
confidentiality footer, slide numbers, layout backgrounds) so slides paste
into company decks with zero drift. 10 x 5.625 in canvas — the Google Slides
default, so files round-trip through Slides cleanly.

Profiles
    master    the Sept-2026 company master deck ("Parspec Overview"):
              pure black, warm glow layouts, ALL-CAPS light titles with an accent
              phrase, numbered cards, panel rows, umber layer fills.  DEFAULT.
    townhall  the Q2-26 town hall deck: left-bar cards, bold mixed-case titles,
              eyebrow labels, numbered section strips.

Templates are company-internal and are NOT shipped in this repo. Build one
from a deck export:

    python3 scripts/make_template.py --profile master "<master deck export>.pptx"

    from assets.house import HouseDeck
    deck = HouseDeck()                      # profile="master"
    deck.cover(["Quarterly", "Update"], date="[Date]", byline="Presenter, Title")
    s = deck.content("Where we stand", accent="this quarter")
    deck.save("deck.pptx")
"""

from __future__ import annotations

import copy
import io
import os
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches

from . import shapes as sh
from .brand import Brand, load as load_brand
from .shapes import Palette

TEMPLATE_DIR = Path.home() / "Documents/Claude/Templates"
W, H = 10.0, 5.625


@dataclass(frozen=True)
class Profile:
    name: str
    template: Path
    env: str
    layouts: dict                  # semantic -> layout name in the kept master
    title_size: float
    title_bold: bool
    title_caps: bool
    x0: float = 0.44               # text edge: title placeholder x + its 0.1" inset
    xw: float = 9.12
    top: float = 1.42              # content band
    bottom: float = 4.95
    notes: dict = field(default_factory=dict)

    @property
    def required(self):
        return set(self.layouts.values())


PROFILES = {
    "master": Profile(
        name="master",
        template=TEMPLATE_DIR / "parspec_master_template.pptx",
        env="PARSPEC_PPTX_MASTER_TEMPLATE",
        layouts={
            "cover": "TITLE_ONLY_1_1",   # black; logo + title drawn by cover()
            "divider": "CUSTOM_5",       # glow top-right + bottom-left, 33pt title
            "content": "CUSTOM_6_2",     # full warm glow + title
            "plain": "CUSTOM_2",         # black + title (no glow)
        },
        title_size=25, title_bold=False, title_caps=True,
        top=1.35, bottom=5.05,
        notes={"white_layouts": ["DEFAULT"], "double_footer_layouts": ["BLANK"]},
    ),
    "townhall": Profile(
        name="townhall",
        template=TEMPLATE_DIR / "parspec_townhall_template.pptx",
        env="PARSPEC_PPTX_TEMPLATE",
        layouts={
            "cover": "CUSTOM_3",                   # logo, title, 2 subtitles, laptop render
            "divider": "CUSTOM_3_1",               # logo, title, subtitle
            "content": "TITLE_AND_TWO_COLUMNS_1",  # left-aligned title
            "plain": "TITLE_AND_TWO_COLUMNS_1",
            "blank": "CUSTOM_2",                   # dark, footer only
        },
        title_size=24, title_bold=True, title_caps=False,
        notes={"white_layouts": ["DEFAULT"]},
    ),
}
DEFAULT_PROFILE = "master"


def template_path(profile: Profile, path: str | Path | None = None) -> Path:
    p = Path(path or os.environ.get(profile.env) or profile.template).expanduser()
    if not p.exists():
        raise FileNotFoundError(
            f"{profile.name} template not found at {p}. Export the latest {profile.name} deck "
            f"as .pptx and run: scripts/make_template.py --profile {profile.name} <file> "
            f"(or set ${profile.env})."
        )
    return p


def strip_to_template(prs: Presentation, keep_master: int) -> Presentation:
    """Delete every slide and every master except `keep_master`, in place.

    python-pptx only writes parts reachable from relationships, so dropped
    slides, masters, layouts and their media vanish on save.
    """
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    keep = prs.slide_masters[keep_master].part
    mids = prs.part._element.find(qn("p:sldMasterIdLst"))
    for mid in list(mids):
        rid = mid.get(qn("r:id"))
        if prs.part.related_part(rid) is not keep:
            prs.part.drop_rel(rid)
            mids.remove(mid)
    return prs


def pick_master(prs: Presentation, profile: Profile) -> int | None:
    for i, m in enumerate(prs.slide_masters):
        if profile.required <= {lay.name for lay in m.slide_layouts}:
            return i
    return None


class HouseDeck:
    """Deck builder bound to a house template profile and the brand palette."""

    def __init__(self, profile: str = DEFAULT_PROFILE, template: str | Path | None = None,
                 brand: Brand | None = None):
        self.p = PROFILES[profile]
        self.brand = brand or load_brand()
        self.pal = Palette.from_brand(self.brand)
        self.prs = Presentation(str(template_path(self.p, template)))
        if len(self.prs.slides) or len(self.prs.slide_masters) > 1:  # un-stripped export
            idx = pick_master(self.prs, self.p)
            if idx is None:
                raise ValueError(f"no master in template has {sorted(self.p.required)}")
            strip_to_template(self.prs, keep_master=idx)
        self.x0, self.xw, self.top, self.bottom = self.p.x0, self.p.xw, self.p.top, self.p.bottom

    # ── plumbing ─────────────────────────────────────────────────────
    @property
    def master(self):
        return self.prs.slide_masters[0]

    def layout(self, kind: str):
        name = self.p.layouts.get(kind, kind)
        for lay in self.master.slide_layouts:
            if lay.name == name:
                return lay
        raise KeyError(f"layout {name!r} not in {self.p.name} template")

    def slide(self, kind: str, keep=(0,)):
        """New slide on a layout; drops unused placeholders; always carries a page number."""
        lay = self.layout(kind)
        s = self.prs.slides.add_slide(lay)
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx not in keep:
                ph._element.getparent().remove(ph._element)
        # add_slide never copies the slide-number field; glow layouts lack one entirely
        src = [ph for ph in lay.placeholders
               if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER] or \
              [ph for ph in self.master.placeholders
               if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER]
        if src:
            s.shapes._spTree.append(copy.deepcopy(src[0]._element))
        return s

    def logo(self, s, x=0.45, y=0.46, w=1.5):
        """Full-resolution PARSPEC wordmark, taken from the master footer."""
        pics = [p for p in self.master.shapes if p.shape_type == 13]
        if not pics:
            return None
        pic = max(pics, key=lambda p: p.image.size[0])
        iw, ih = pic.image.size
        return s.shapes.add_picture(io.BytesIO(pic.image.blob), Inches(x), Inches(y),
                                    Inches(w), Inches(w * ih / iw))

    @staticmethod
    def place(ph, dx=0.0, dy=0.0, *, x=None, y=None, w=None, h=None):
        """Move a placeholder. Always writes all four values: setting only .top
        on an inheriting placeholder writes an xfrm with left=0."""
        l, t, wd, ht = ph.left, ph.top, ph.width, ph.height
        ph.left = Inches(x) if x is not None else l + Inches(dx)
        ph.top = Inches(y) if y is not None else t + Inches(dy)
        ph.width = Inches(w) if w is not None else wd
        ph.height = Inches(h) if h is not None else ht

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        self.prs.save(path)
        return path

    @staticmethod
    def notes(s, text: str):
        s.notes_slide.notes_text_frame.text = text.strip()

    def placeholder(self, t: str):
        return sh.placeholder(t, self.pal)

    def _case(self, t: str) -> str:
        return t.upper() if self.p.title_caps else t

    # ── chrome ───────────────────────────────────────────────────────
    def eyebrow(self, s, label: str):
        sh.label(s, self.x0, 0.26, 4.0, label, self.pal)

    def title(self, s, text: str, accent: str | None = None, *, top=None):
        """Title placeholder; `accent` is an optional trailing phrase in brand orange."""
        ph = s.shapes.title
        self.place(ph, x=0.34, y=top if top is not None else (0.49 if self.p.title_caps else 0.36),
                   w=9.32, h=0.63)
        runs = [(self._case(text) + (" " if accent else ""), {"color": self.pal.text})]
        if accent:
            runs.append((self._case(accent), {"color": self.pal.accent}))
        sh.text(s, 0, 0, 0, 0, [runs], size=self.p.title_size, bold=self.p.title_bold,
                target=ph)
        tf = ph.text_frame  # restore the layout's 0.1" insets that text() zeroed
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.1)

    def subtitle(self, s, text: str, y=None):
        y = y if y is not None else (1.04 if self.p.title_caps else 0.98)
        sh.text(s, self.x0, y, self.xw, 0.26, text, size=11, color=self.pal.muted)

    # ── slide recipes ────────────────────────────────────────────────
    def content(self, title: str, *, accent=None, eyebrow=None, subtitle=None, glow=True):
        """Titled content slide. master: glow=False gives the plain black layout."""
        s = self.slide("content" if glow else "plain")
        if eyebrow:
            self.eyebrow(s, eyebrow)
        self.title(s, title, accent)
        if subtitle:
            self.subtitle(s, subtitle)
        return s

    def cover(self, title_lines, subtitle: str = "", byline="", *, date="",
              byline_label="Presented by", eyebrow: str | None = None):
        """byline / date: str or runs (use self.placeholder('[Date]') for unknowns)."""
        if self.p.name == "townhall":
            return self._cover_townhall(title_lines, subtitle, byline, date, eyebrow)
        pal = self.pal
        s = self.slide("cover", keep=())
        self.logo(s)
        if eyebrow:
            sh.label(s, 0.52, 1.2, 6, eyebrow, pal, size=11)
        sh.text(s, 0.52, 1.45, 7.5, 1.75, list(title_lines), size=40, bold=True,
                color=pal.text, anchor="b", line_spacing=1.0)
        sh.rule(s, 0.52, 3.3, 2.25, pal.accent, h=0.02)
        y = 3.45
        if subtitle:
            sh.text(s, 0.52, y, 7, 0.3, subtitle, size=14, color=pal.text)
            y += 0.4
        if date:
            sh.text(s, 0.52, y, 7, 0.25, [sh.runs(date)], size=11, color=pal.subtle)
            y += 0.38
        if byline:
            sh.text(s, 0.52, y, 7, 0.22, byline_label, size=11, bold=True, color=pal.accent)
            sh.text(s, 0.52, y + 0.22, 7, 0.25, [sh.runs(byline)], size=11, color=pal.text)
        sh.arcs(s, pal, x=8.67, y=4.41, d=2.18, corner="br")
        return s

    def _cover_townhall(self, title_lines, subtitle, byline, date, eyebrow):
        s = self.slide("cover", keep=(0, 1, 2))
        ph = {p.placeholder_format.idx: p for p in s.placeholders}
        if eyebrow:
            sh.shape(s, MSO_SHAPE.RECTANGLE, 0.43, 0.98, 0.27, 0.012, fill=self.pal.accent)
            sh.text(s, 0.79, 0.92, 3.5, 0.16, eyebrow.upper(), size=8.5, bold=True,
                    color=self.pal.accent)
        sh.text(s, 0, 0, 0, 0, list(title_lines), size=38, bold=True, anchor="b",
                color=self.pal.text, target=ph[0])
        sh.text(s, 0, 0, 0, 0, subtitle, size=14, color=self.pal.text, target=ph[1])
        line = []
        for part in (date, byline):
            if part:
                if line:
                    line.append(("  ·  ", {}))
                line += sh.runs(part)
        sh.text(s, 0, 0, 0, 0, [line or [("", {})]], size=11, color=self.pal.muted, target=ph[2])
        ph[0].text_frame.margin_left = Inches(0.1)
        for p in (ph[1], ph[2]):
            p.text_frame.margin_left = Inches(0.1)
            self.place(p, dx=-0.15, dy=0.12)  # the layout indents body text
        return s

    def divider(self, heading: str, minutes: str | None = None, *, accent=None, kicker=None,
                strip=None, active=None, tagline: str | None = None, footnote=None, size=None):
        """Section divider.

        master: glow layout, 33pt caps heading (+ accent phrase), optional kicker
        line above ('PARTNER WITH US'), tagline/minutes/footnote below.
        townhall: logo, arcs, big heading, optional numbered strip.
        """
        pal = self.pal
        if self.p.name == "master":
            s = self.slide("divider")
            ph = s.shapes.title
            runs = [(self._case(heading) + (" " if accent else ""), {"color": pal.text})]
            if accent:
                runs.append((self._case(accent), {"color": pal.accent}))
            sh.text(s, 0, 0, 0, 0, [runs], size=size or 33, target=ph, line_spacing=1.05)
            ph.text_frame.margin_left = Inches(0.1)
            if kicker:
                sh.text(s, 0.72, 1.78, 6, 0.25, kicker.upper(), size=12, color=pal.text)
            y = 3.62
            for part, kw in ((tagline, dict(size=14, bold=True, color=pal.accent)),
                             (minutes, dict(size=12, color=pal.subtle))):
                if part:
                    sh.text(s, 0.72, y, 7, 0.3, part, **kw)
                    y += 0.38
            if footnote:
                sh.text(s, 0.72, y, 7, 0.25, [sh.runs(footnote)], size=11, color=pal.muted)
            if strip:
                sh.section_strip(s, strip, pal, active=active, y=4.55)
            return s

        s = self.slide("divider", keep=(0, 1))
        sh.arcs(s, pal)
        ph = {p.placeholder_format.idx: p for p in s.placeholders}
        head = [(heading + (" " if accent else ""), {})] + ([(accent, {"color": pal.accent})] if accent else [])
        sh.text(s, 0, 0, 0, 0, [head], size=30, bold=True, anchor="b",
                color=pal.text, target=ph[0])
        sh.text(s, 0, 0, 0, 0, minutes or "", size=14, color=pal.text, target=ph[1])
        for p in ph.values():
            p.text_frame.margin_left = Inches(0.1)
        if tagline:
            sh.text(s, self.x0, 3.72, 6.0, 0.3, tagline, size=14, bold=True, color=pal.accent)
        if footnote:
            sh.text(s, self.x0, 4.12, 6.0, 0.25, [sh.runs(footnote)], size=11, color=pal.muted)
        if strip:
            sh.section_strip(s, strip, pal, active=active)
        return s

    def agenda(self, items, *, heading=("TODAY'S", "AGENDA"), blurb="", eyebrow=None,
               highlight=None):
        """items: [(title, descriptor, minutes)]. Numbered — an agenda is a sequence.

        highlight: index of the row that should read as the climax (master: accent title).
        """
        pal = self.pal
        if self.p.name == "master":
            s = self.slide("plain", keep=())
            sh.text(s, 0.45, 1.9, 2.6, 1.2, list(heading), size=30, font="Montserrat Light",
                    color=pal.text, line_spacing=1.0)
            if blurb:
                sh.text(s, 0.45, 3.2, 2.4, 0.5, blurb, size=10, color=pal.muted)
            sh.shape(s, MSO_SHAPE.RECTANGLE, 3.35, 1.25, 0.01, 3.4, fill=pal.card_line)
            x, y = 3.75, 1.25
            rh = min(0.62, 3.4 / len(items))
            for i, (t, d, m) in enumerate(items):
                yy = y + i * rh
                sh.text(s, x, yy, 0.6, rh, f"{i + 1:02d}", size=20, bold=True,
                        color=pal.accent, anchor="m")
                sh.text(s, x + 0.75, yy, 3.9, rh,
                        [[(d.upper(), {"size": 9, "color": pal.subtle})],
                         [(t, {"bold": i == highlight})]],
                        size=15, color=pal.accent if i == highlight else pal.text, anchor="m")
                sh.text(s, x + 4.6, yy, 0.9, rh, m, size=11, color=pal.accent, anchor="m",
                        align="r")
                sh.rule(s, x, yy + rh - 0.005, 5.5, pal.card_line)
            return s

        s = self.slide("blank", keep=())
        if eyebrow:
            self.eyebrow(s, eyebrow)
        sh.text(s, self.x0, 1.95, 2.5, 0.9, list(heading), size=26, bold=True, color=pal.accent)
        if blurb:
            sh.text(s, self.x0, 2.95, 2.5, 0.4, blurb, size=9, color=pal.muted)
        x, y = 3.25, 1.02
        rh = min(0.6, 3.75 / len(items))
        for i, (t, d, m) in enumerate(items):
            yy = y + i * rh
            sh.text(s, x, yy, 0.5, rh, f"{i + 1:02d}", size=15, bold=True,
                    color=pal.accent, anchor="m")
            sh.text(s, x + 0.55, yy, 2.55, rh, t, size=13, bold=True, color=pal.text,
                    anchor="m")
            sh.text(s, x + 3.15, yy, 2.25, rh, d, size=9, italic=True, color=pal.muted,
                    anchor="m")
            sh.text(s, x + 5.45, yy, 0.86, rh, m, size=9, bold=True, color=pal.accent,
                    anchor="m", align="r")
            if i < len(items) - 1:
                sh.rule(s, x, yy + rh - 0.005, 6.31, pal.card_line)
        return s
