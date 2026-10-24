"""Editable Series B story archetypes. Caller-owned data; no company content.

All geometry is on the investor 1280 x 720 canvas. These builders do not infer
product status, commercial rights, missing outcomes or cross-meter conversions.
"""
from decimal import Decimal, ROUND_HALF_UP
import math
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

ILLUSTRATIVE = 'ILLUSTRATIVE · DUMMY NUMBERS'
DISCLAIMER = 'Illustrative — dummy numbers, not Parspec data'
STATUSES = ('Live', 'Demo environment', 'Prototype', 'Roadmap')


def number(value, places=1):
    return f'{Decimal(str(value)).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP):,.{places}f}'


def illustrative_chip(deck, slide):
    i = deck.investor
    i.rect(slide, 844, 32, 380, 32, fill=None, line='accent', lw=1).name = 'incubation:outline'
    i.text(slide, 852, 39, 364, 22, ILLUSTRATIVE, size=14, align='c', bold=True).name = 'incubation:chip'
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.startswith('©'):
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.text = DISCLAIMER
                    return


def _slide(deck, sid, title, subtitle='', incubation=False, notes=''):
    s = deck.slide('plain', keep=())
    i = deck.investor
    if '\n' not in title and len(title)>52:
        words=title.split()
        middle=min(range(1,len(words)),key=lambda j:abs(len(' '.join(words[:j]))-len(title)/2))
        title=' '.join(words[:middle])+'\n'+' '.join(words[middle:])
    i.header(s, sid + (' · INCUBATION' if incubation and sid!='INCUBATION' else ''), title, subtitle)
    if incubation:
        illustrative_chip(deck, s)
    deck.notes(s, f'Slide: {sid}\n' + (DISCLAIMER+'\n' if incubation else '') + notes)
    return s


def _table(deck, s, headers, rows, *, x=56, y=232, w=1168, h=400, widths=None, numeric_cols=(), font_px=18):
    return deck.charts.fin_table(s, headers, rows, x=x, y=y, w=w, h=h,
                                widths=widths, numeric_cols=numeric_cols, font_px=font_px)


def _connector(deck, s, x1,y1,x2,y2, *, arrow=True, name='workflow:connector'):
    ppi=deck.investor.ppi
    line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[Inches(v/ppi) for v in (x1,y1,x2,y2)])
    line.name=name; line.line.color.rgb=deck.investor.color('muted'); line.line.width=deck.charts.pt(1)
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');line._element.spPr.get_or_add_ln().append(end)
    return line


def _status(deck,s,x,y,w,label):
    if label not in STATUSES:
        raise ValueError('unsupported capability status: '+label)
    deck.investor.rect(s,x,y,w,30,fill='band',line='line')
    deck.investor.text(s,x+8,y+5,w-16,24,label,size=14,bold=True,align='c')


def _columns(deck,s,categories,values, *,x,y,w,h,name,maximum=None,percent=False):
    ch=deck.charts; data=CategoryChartData();data.categories=categories
    for item in values: data.add_series(item['name'],item['values'])
    shape=ch._new(s,XL_CHART_TYPE.COLUMN_CLUSTERED,data,x,y,w,h,name,font_px=16)
    for k,ser in enumerate(shape.chart.series): ch._paint(ser,'chart_2' if k==0 else 'chart_4')
    shape.chart.plots[0].gap_width=70
    ch._labels(shape.chart.plots[0],fmt='0.0',font_px=16)
    ch._axes(shape.chart,font_px=16,minimum=0,maximum=maximum,number_format='0')
    return shape


def incubation_divider(deck, data=None):
    s=_slide(deck,'INCUBATION','Incubation · Economics we will prove',incubation=True)
    lines=[
        'D01  Approved market scope, source version and eligible buyers',
        'D02 + D05  Matched earned revenue and full delivery cost, tied to books',
        'D03  Task and attempt ledger, acceptance, retries and human review',
        'D04  Customer baseline, attributable benefit and quality evaluation',
        'D06  Tested willingness to pay, pilot meter and refund rules',
    ]
    for j,line in enumerate(lines):
        deck.investor.text(s,56,240+j*70,1168,52,line,size=24)
    return s


def wedge_workflow(deck, data):
    s=_slide(deck,'IN-20',data.get('title','The paid wedge and the next supervised workflow'),data.get('subtitle','Generic placeholder scope. Replace with a dated product inventory.'),notes=data.get('notes',''))
    columns=data['columns']
    rows=[[label]+[c[key] for c in columns] for label,key in [('Buyer / job','buyer_job'),('Status','status'),('Reusable context','context'),('Next dependency','dependency')]]
    _table(deck,s,['Workflow']+[c['name'] for c in columns],rows,h=390,widths=[.16]+[.84/len(columns)]*len(columns),font_px=18)
    for j in range(len(columns)-1):
        boundary=56+1168*.16+(j+1)*1168*.84/len(columns)
        _connector(deck,s,boundary-64,271,boundary-24,271)
    return s


def agent_swimlanes(deck, data):
    s=_slide(deck,'IN-21',data.get('title','Agent preparation and human controls'),data.get('subtitle','Generic placeholder flow with explicit capability stages.'),notes=data.get('notes',''))
    i=deck.investor
    for label,key,yy,height in [('TRIGGER','triggers',226,82),('AGENT STEPS','agent_steps',320,176),('HUMAN CONTROLS','human_controls',508,122)]:
        items=data[key];cols=min(3,len(items));cw=940/cols
        row_count=(len(items)+cols-1)//cols;rh=height/row_count
        i.text(s,56,yy+8,190,44,label,size=18,bold=True,color='muted')
        i.rule(s,56,yy+height,1168)
        for j,item in enumerate(items):
            x=266+(j%cols)*cw;iy=yy+(j//cols)*rh
            has_status=bool(item.get('status'))
            i.text(s,x,iy,cw-28,min(rh-32,64) if has_status else rh-2,item['text'],size=14 if len(items)>3 else 18,bold=True)
            if has_status: _status(deck,s,x,iy+min(rh-33,62),min(cw-28,220),item['status'])
            if j%cols<cols-1: _connector(deck,s,x+cw-25,iy+20,x+cw-5,iy+20)
    _connector(deck,s,1208,308,1208,320)
    _connector(deck,s,1208,496,1208,508)
    i.text(s,56,642,1168,28,data.get('audit_path','Audit output: task ID, attempts, approvals and exception log'),size=16,color='muted')
    return s


def learning_ledger(deck, data):
    s=_slide(deck,'IN-22',data.get('title','Design-partner reviews shape the workflow'),data.get('subtitle',data.get('partner_label','[Design partner]')),notes=data.get('notes',''))
    _table(deck,s,['Customer request','Product response / spec','Observed status','Next test'],[[r[k] for k in ('request','response','status','next_test')] for r in data['rows']],h=330,widths=[.24,.29,.19,.28],font_px=16)
    stages=data['chronology'];cw=1168/len(stages)
    for j,stage in enumerate(stages):
        deck.investor.rect(s,56+j*cw,572,cw-4,94,fill='band')
        deck.investor.text(s,68+j*cw,580,cw-26,84,stage,size=14)
        if j<len(stages)-1: _connector(deck,s,56+(j+1)*cw-16,618,56+(j+1)*cw+10,618)
    return s


def reachable_revenue(deck, data):
    totals=[Decimal(str(r['parents']))*(Decimal(str(r['platform_fee']))+Decimal(str(r['units_per_parent']))*Decimal(str(r['net_unit_price']))) / 1000000 for r in data['rows']]
    total=sum(totals)
    s=_slide(deck,'IN-23',f'${number(total)}M annual vendor revenue scenario','Annual US$M. Separate serviceable and expansion cases.',True)
    rows=[[r['segment'],str(r['parents']),number(r['units_per_parent'],0), '$'+number(r['net_unit_price'],2),'$'+number(r['platform_fee'],0),number(v)] for r,v in zip(data['rows'],totals)]
    rows.append(['Total','','','','',number(total)])
    _table(deck,s,['Segment','Eligible\nparents','Units /\nparent / yr','Net $ /\nunit','Platform\n$ / yr','Revenue\n$M / yr'],rows,w=700,h=386,widths=[.28,.12,.16,.14,.16,.14],numeric_cols=[1,2,3,4,5],font_px=14)
    _columns(deck,s,['Serviceable\ncase','Expansion\nscenario'],[dict(name='Annual revenue',values=[float(total),data['expansion_scenario_m']])],x=780,y=218,w=444,h=370,name='reachable_revenue')
    point=s.shapes[-1].chart.series[0].points[1]
    point.format.fill.solid();point.format.fill.fore_color.rgb=deck.investor.color('chart_4')
    point.format.line.color.rgb=deck.investor.color('chart_4')
    deck.investor.text(s,56,624,1168,40,'Unit: '+data['unit']+'. Revenue = parents × (platform fee + incremental units × net unit price).',size=16,color='muted')
    return s


def _waterfall(deck, slide, labels, values, **kwargs):
    shape=deck.charts.waterfall(slide,labels,values,**kwargs)
    totals={j%len(values) for j in kwargs.get('total_indices',(0,-1))}
    movement=shape.chart.series[1]
    for j in totals:
        point=movement.points[j]
        point.format.fill.solid();point.format.fill.fore_color.rgb=deck.investor.color('chart_2')
        point.format.line.color.rgb=deck.investor.color('chart_2')
    return shape


def labor_waterfall(deck, data):
    s=_slide(deck,'IN-24','Labor wallet and vendor revenue scenario',data['unit'],True)
    _waterfall(deck,s,[r['label'].replace(' ','\n',1) for r in data['steps']],[r['value'] for r in data['steps']],total_indices=[j for j,r in enumerate(data['steps']) if r['type']=='total'],x=56,y=220,w=1168,h=260,font_px=16,number_format='0',label_above=True)
    _table(deck,s,['Factor','Applied once','Evidence'],[
        ['Technical coverage',data['steps'][1]['label'],'Dummy assumption'],['Benefit realization',data['steps'][2]['label'],'Dummy assumption'],['Vendor capture',number(data['steps'][-1]['value']/data['steps'][3]['value']*100)+'% of benefit','Dummy scenario']],y=522,h=136,widths=[.32,.38,.30],font_px=16)
    return s


def earned_expansion(deck, data):
    s=_slide(deck,'IN-25','Each workflow earns a separate expansion case','Annual revenue increments require separate pricing and overlap checks.',True)
    headers=['Workflow','Buyer','Incremental unit','Dependency','Overlap / pricing proof','Readiness']
    def row(r): return [r['workflow'],r['buyer'],r['unit'],r['dependency'],r.get('overlap_pricing','Overlap removal and price test pending'),r['readiness']]
    _table(deck,s,headers,[row(r) for r in data[:-1]],h=280,widths=[.23,.15,.16,.15,.18,.13],font_px=16)
    deck.investor.text(s,56,532,1168,28,'SEPARATE MARKETPLACE ENGINE · EXCLUDED FROM EVERY TOTAL',size=18,bold=True)
    _table(deck,s,headers,[row(data[-1])],y=570,h=96,widths=[.23,.15,.16,.15,.18,.13],font_px=14)
    return s


def pilot_package(deck, data):
    p=data['pilot']
    s=_slide(deck,'IN-26','Proposed pilot package: commitment plus accepted work','Commercial terms are hypothetical. Existing meters remain distinct.',True)
    rows=[['Pilot / quoting team',p['unit'],number(p['included_units'],0)+' included',f"${number(p['base_commitment'],0)} / {p['term']}",f"${number(p['overage_price'],2)} overage. "+p['refund_rule'],'Proposed']]
    for r in data['existing_meters']: rows.append([r['meter'],r['billable_event'],'Per contract','Per contract','Per contract',r['status']])
    _table(deck,s,['Buyer / job','Billable event','Included volume','Price basis','Overage / refund','Status'],rows,h=370,widths=[.18,.18,.13,.15,.23,.13],font_px=18)
    deck.investor.text(s,56,620,1168,44,'Acceptance, duplicate and reopen rules must prevent double billing.',size=20,color='muted')
    return s


def output_economics(deck, data):
    p=data['per_accepted_output'];t=data['table']
    s=_slide(deck,'IN-27','Value, price and full cost per accepted output','US$ / accepted output. Benefit and price overlap economically.',True)
    scale=max(1,math.ceil(max(p.values())*1.1))
    chart=_columns(deck,s,['Customer\nbenefit','Allocated\nnet price','Full delivery\ncost'],[dict(name='US$ / accepted output',values=[p[k] for k in ('customer_benefit','net_price','full_cost')])],x=56,y=214,w=760,h=310,name='output_economics',maximum=scale)
    area=chart.chart._chartSpace.chart.plotArea
    layout=area.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout')
    if layout is None:layout=OxmlElement('c:layout');area.insert(0,layout)
    manual=OxmlElement('c:manualLayout')
    for name,val in [('layoutTarget','inner'),('xMode','edge'),('yMode','edge'),('wMode','factor'),('hMode','factor'),('x',.08),('y',.05),('w',.88),('h',.76)]:
        e=OxmlElement('c:'+name);e.set('val',str(val));manual.append(e)
    layout.append(manual)
    for j,(top,bottom,label) in enumerate([(p['customer_benefit'],p['net_price'],'Customer retained\nbefore adoption cost'),(p['net_price'],p['full_cost'],'Vendor gross profit')]):
        x=850; y1=214+310*.05+310*.76*(1-top/scale);y2=214+310*.05+310*.76*(1-bottom/scale)
        _connector(deck,s,x,y1,x,y2,arrow=False,name='economics:bracket')
        for yy in (y1,y2): _connector(deck,s,x-8,yy,x+8,yy,arrow=False,name='economics:bracket-tick')
        deck.investor.text(s,x+24,(y1+y2)/2-28 if j==0 else y1,340,84,f'${number(top-bottom,2)}  '+label,size=22,bold=True)
    _table(deck,s,['Review min / output','Acceptance','Earned revenue','Accepted outputs','Period'],[[number(t['review_minutes']),number(t['acceptance_rate_pct'])+'%','$'+number(t['earned_revenue_k'])+'K',number(t['accepted_outputs'],0),t['period']]],y=558,h=100,widths=[.20,.15,.18,.18,.29],font_px=18)
    return s


def delivery_payback(deck, data):
    steps=data['waterfall_k'];p=data['payback'];running=0;cumulative=[]
    for v in p['monthly_gp_k']: running+=v;cumulative.append(running)
    month=next((j+1 for j,v in enumerate(cumulative) if v>=p['deploy_cost_k']),None)
    s=_slide(deck,'IN-28',f'${number(sum(r["value"] for r in steps[:-1]),0)}K gross profit after delivery costs','US$K. Cohort period: UNKNOWN. Illustrative months since go-live.',True)
    labels=['Earned\nrevenue','Inference','Tools /\ndata','Hosting','Human QA /\nrework','Support /\nimpl.','Other\ndelivery','Gross\nprofit']
    _waterfall(deck,s,labels,[r['value'] for r in steps],x=56,y=244,w=710,h=342,font_px=14,number_format='0',label_above=True)
    deck.investor.text(s,800,204,424,40,'CUMULATIVE GP AND DEPLOYMENT COST',size=16,bold=True)
    chart=deck.charts.line_trend(s,[str(j+1) for j in range(len(cumulative))],[dict(name='Cumulative GP',values=cumulative,color='chart_2'),dict(name='Acquisition / deployment cost',values=[p['deploy_cost_k']]*len(cumulative),color='muted')],x=800,y=248,w=424,h=300,font_px=14,number_format='0',legend=False)
    chart.chart.plots[0].has_data_labels=False
    # Series-level defaults must also be explicit: importers otherwise label
    # every point when the single native payback label is added below.
    for series in chart.chart.series:
        labels_xml=OxmlElement('c:dLbls')
        for option in ('showLegendKey','showVal','showCatName','showSerName','showPercent','showBubbleSize'):
            flag=OxmlElement('c:'+option);flag.set('val','0');labels_xml.append(flag)
        series._element.insert(list(series._element).index(series._element.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}cat')),labels_xml)
    skip=OxmlElement('c:tickLblSkip');skip.set('val','3');chart.chart.category_axis._element.append(skip)
    if month:
        deck.charts._custom_label(chart.chart.series[0],month-1,'Month '+str(month),font_px=14,position='t')
    deck.investor.text(s,800,566,424,76,'Payback: '+('month '+str(month) if month else 'UNKNOWN')+'\nGP crosses $'+number(p['deploy_cost_k'],0)+'K cost',size=20,bold=True)
    deck.investor.text(s,56,622,710,36,'All attempts, failed runs and human delivery included.',size=18,color='muted')
    return s


def deployment_cohorts(deck, data):
    s=_slide(deck,'IN-29','Paid and live deployment cohorts','All signed customers remain in the denominator. Unmatured renewals are unknown.',True)
    _table(deck,s,['Fiscal cohort','Signed','Paid','Live','Renewed','Implementation\nmedian days','Custom\nhours'],[[r['cohort'],str(r['signed']),str(r['paid']),str(r['live']),'UNKNOWN' if r['renewed'] is None else str(r['renewed']),str(r['impl_days_median']),str(r['custom_hours'])] for r in data],h=350,widths=[.21,.09,.09,.09,.16,.19,.17],numeric_cols=[1,2,3,4,5,6],font_px=18)
    return s


def context_evidence(deck, data):
    s=_slide(deck,'IN-30','Permissioned context and held-out evaluations','Each result needs a defined metric, access right and reproducible comparison.',True)
    _table(deck,s,['Context / integration','Right / access','Evaluation','Baseline','With context'],[[r['context'],r['right'],r['eval'],number(r['baseline_pct'])+'%',number(r['with_context_pct'])+'%'] for r in data],h=390,widths=[.25,.24,.27,.12,.12],numeric_cols=[3,4],font_px=20)
    return s


def quality_leverage(deck, data):
    s=_slide(deck,'IN-31','Final outcomes, review burden and delivery cost','Fixed workload mix and QA rubric required. Observation windows remain unknown. TBD = not measured.',True)
    cats=[r['release'] for r in data];ch=deck.charts
    series=[dict(name=label,values=[r[key]/100 for r in data],color=('chart_2','chart_4','chart_3')[j]) for j,(key,label) in enumerate([('accepted_no_review_pct','Accepted, no review'),('accepted_with_review_pct','Accepted, reviewed'),('failed_pct','Failed / unresolved')])]
    shape=ch.stacked_bar(s,cats,series,x=56,y=218,w=558,h=255,horizontal=False,font_px=14,number_format='0%',maximum=1)
    # A genuine percentStacked native chart, not a standard stack that happens to sum to one.
    shape.chart._chartSpace.xpath('.//c:barChart/c:grouping')[0].set('val','percentStacked')
    rows=[[label]+[str(r.get(key,'TBD')) for r in data] for label,key in [('Eligible','eligible'),('Attempted','attempted'),('Accepted','accepted'),('Defect / reopen %','defect_reopen_pct'),('Observation days','observation_days')]]
    _table(deck,s,['Counts / period']+cats,rows,x=650,y=218,w=574,h=245,widths=[.36,.16,.16,.16,.16],numeric_cols=[1,2,3,4],font_px=14)
    for j,(key,label) in enumerate([('review_min_per_output','REVIEW MINUTES / ACCEPTED OUTPUT'),('full_cost_per_output','FULL US$ COST / ACCEPTED OUTPUT')]):
        xx=56+j*594
        deck.investor.text(s,xx,490,574,28,label,size=16,bold=True)
        ch.line_trend(s,cats,[dict(name=label,values=[r[key] for r in data],color='chart_2')],x=xx,y=526,w=574,h=136,font_px=14,number_format='0.0')
    return s


BUILDERS = {23:reachable_revenue,24:labor_waterfall,25:earned_expansion,26:pilot_package,
            27:output_economics,28:delivery_payback,29:deployment_cohorts,30:context_evidence,31:quality_leverage}
