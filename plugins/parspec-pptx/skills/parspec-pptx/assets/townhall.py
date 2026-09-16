"""
parspec-pptx — Town Hall register.

Builds decks on the company's own town hall template (the dark Google Slides
master: PARSPEC logo, confidentiality footer, slide numbers) so slides paste
into all-hands / town hall decks with zero drift. 10 x 5.625 in canvas —
the Google Slides default, so the file round-trips through Slides cleanly.

The template is company-internal and is NOT shipped in this repo. Build it
from any town hall export:

    python3 scripts/make_template.py "Parspec Townhall - <quarter>.pptx"

It lands at DEFAULT_TEMPLATE; override with $PARSPEC_PPTX_TEMPLATE.

    from assets.townhall import TownHall
    th = TownHall()
    th.cover(["Quarterly", "Town Hall"], "Company update", "Presenter, Title")
    s = th.content("Where We Stand", accent="This Quarter", eyebrow="Update",
                   subtitle="One-line takeaway.")
    th.save("deck.pptx")
"""

from __future__ import annotations

import copy
import os
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .brand import Brand, load as load_brand
from . import shapes as sh
from .shapes import FONT, Palette

TEMPLATE_ENV = "PARSPEC_PPTX_TEMPLATE"
DEFAULT_TEMPLATE = Path.home() / "Documents/Claude/Templates/parspec_townhall_template.pptx"

# Semantic name -> layout name inside the town hall master (Google Slides export).
LAYOUTS = {
    "cover": "CUSTOM_3",                   # logo, title, 2 subtitles, product laptop render
    "divider": "CUSTOM_3_1",               # logo, title, subtitle
    "content": "TITLE_AND_TWO_COLUMNS_1",  # left-aligned title
    "blank": "CUSTOM_2",                   # dark canvas, footer only
}
# The master's DEFAULT layout is WHITE — never build on it.

W, H = 10.0, 5.625
X0, XW = 0.44, 9.12            # text edge (title placeholder x + its 0.1" inset), usable width
TOP, BOTTOM = 1.42, 4.95        # content band; the template footer starts at ~5.2


def template_path(path: str | Path | None = None) -> Path:
    p = Path(path or os.environ.get(TEMPLATE_ENV) or DEFAULT_TEMPLATE).expanduser()
    if not p.exists():
        raise FileNotFoundError(
            f"Town hall template not found at {p}. Export the latest town hall deck "
            f"as .pptx and run scripts/make_template.py on it, or set ${TEMPLATE_ENV}."
        )
    return p


def strip_to_template(prs: Presentation, keep_master: int) -> Presentation:
    """Delete every slide and every master except `keep_master`, in place.

    python-pptx only writes parts reachable from relationships, so dropped
    slides, masters, layouts and their media vanish on save.
    """
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    keep = prs.slide_masters[keep_master].part
    mids = prs.part._element.find(qn("p:sldMasterIdLst"))
    for mid in list(mids):
        rid = mid.get(qn("r:id"))
        if prs.part.related_part(rid) is not keep:
            prs.part.drop_rel(rid)
            mids.remove(mid)
    return prs


class TownHall:
    """Deck builder bound to the town hall template and a brand palette."""

    def __init__(self, template: str | Path | None = None, brand: Brand | None = None):
        self.brand = brand or load_brand()
        self.pal = Palette.from_brand(self.brand)
        self.prs = Presentation(str(template_path(template)))
        if len(self.prs.slides):  # tolerate an un-stripped export
            strip_to_template(self.prs, keep_master=0)

    # ── plumbing ─────────────────────────────────────────────────────
    def layout(self, kind: str):
        name = LAYOUTS.get(kind, kind)
        for m in self.prs.slide_masters:
            for lay in m.slide_layouts:
                if lay.name == name:
                    return lay
        raise KeyError(f"layout {name!r} not in template")

    def slide(self, kind: str, keep=(0,)):
        """New slide on a layout; drops unused placeholders, keeps the page number."""
        lay = self.layout(kind)
        s = self.prs.slides.add_slide(lay)
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx not in keep:
                ph._element.getparent().remove(ph._element)
        for ph in lay.placeholders:  # add_slide never copies the slide-number field
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                s.shapes._spTree.append(copy.deepcopy(ph._element))
        return s

    @staticmethod
    def place(ph, dx=0.0, dy=0.0, *, x=None, y=None, w=None, h=None):
        """Move a placeholder. Always writes all four values: setting only .top
        on an inheriting placeholder writes an xfrm with left=0."""
        l, t, wd, ht = ph.left, ph.top, ph.width, ph.height
        ph.left = Inches(x) if x is not None else l + Inches(dx)
        ph.top = Inches(y) if y is not None else t + Inches(dy)
        ph.width = Inches(w) if w is not None else wd
        ph.height = Inches(h) if h is not None else ht

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        self.prs.save(path)
        return path

    # ── chrome ───────────────────────────────────────────────────────
    def eyebrow(self, s, label: str):
        sh.label(s, X0, 0.26, 4.0, label, self.pal)

    def title(self, s, text: str, accent: str | None = None, *, size=24, top=0.36):
        """White title; `accent` is an optional trailing orange phrase (house style)."""
        ph = s.shapes.title
        self.place(ph, x=0.34, y=top, w=9.32, h=0.63)
        parts = [(text + (" " if accent else ""), self.pal.text)]
        if accent:
            parts.append((accent, self.pal.accent))
        sh.text(s, 0, 0, 0, 0, [[(t, {"color": c}) for t, c in parts]], size=size,
                bold=True, target=ph)
        tf = ph.text_frame  # restore the layout's 0.1" insets that text() zeroed
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.1)

    def subtitle(self, s, text: str, y=0.98):
        sh.text(s, X0, y, XW, 0.26, text, size=11, color=self.pal.muted)

    @staticmethod
    def notes(s, text: str):
        s.notes_slide.notes_text_frame.text = text.strip()

    def placeholder(self, t: str):
        return sh.placeholder(t, self.pal)

    # ── slide recipes ────────────────────────────────────────────────
    def cover(self, title_lines, subtitle: str, byline, *, eyebrow: str | None = None):
        """byline: str or runs (use self.placeholder('[Date]') for unknowns)."""
        s = self.slide("cover", keep=(0, 1, 2))
        ph = {p.placeholder_format.idx: p for p in s.placeholders}
        if eyebrow:
            sh.shape(s, MSO_SHAPE.RECTANGLE, 0.43, 0.98, 0.27, 0.012, fill=self.pal.accent)
            sh.text(s, 0.79, 0.92, 3.5, 0.16, eyebrow.upper(), size=8.5, bold=True,
                    color=self.pal.accent)
        sh.text(s, 0, 0, 0, 0, list(title_lines), size=38, bold=True, anchor="b",
                color=self.pal.text, target=ph[0])
        sh.text(s, 0, 0, 0, 0, subtitle, size=14, color=self.pal.text, target=ph[1])
        runs = [(byline, {})] if isinstance(byline, str) else byline
        sh.text(s, 0, 0, 0, 0, [runs], size=11, color=self.pal.muted, target=ph[2])
        ph[0].text_frame.margin_left = Inches(0.1)
        for p in (ph[1], ph[2]):
            p.text_frame.margin_left = Inches(0.1)
            self.place(p, dx=-0.15, dy=0.12)  # the layout indents body text
        return s

    def divider(self, heading: str, minutes: str | None = None, *, strip=None,
                active=None, tagline: str | None = None, footnote=None):
        """Section divider: logo, arcs, big heading, optional numbered strip."""
        s = self.slide("divider", keep=(0, 1))
        sh.arcs(s, self.pal)
        ph = {p.placeholder_format.idx: p for p in s.placeholders}
        sh.text(s, 0, 0, 0, 0, heading, size=30, bold=True, anchor="b",
                color=self.pal.text, target=ph[0])
        sh.text(s, 0, 0, 0, 0, minutes or "", size=14, color=self.pal.text, target=ph[1])
        for p in ph.values():
            p.text_frame.margin_left = Inches(0.1)
        if tagline:
            sh.text(s, X0, 3.72, 6.0, 0.3, tagline, size=14, bold=True, color=self.pal.accent)
        if footnote:
            runs = [(footnote, {})] if isinstance(footnote, str) else footnote
            sh.text(s, X0, 4.12, 6.0, 0.25, [runs], size=11, color=self.pal.muted)
        if strip:
            sh.section_strip(s, strip, self.pal, active=active)
        return s

    def agenda(self, items, *, heading=("TODAY'S", "AGENDA"), blurb="", eyebrow=None):
        """items: [(title, descriptor, minutes)]. Numbered — an agenda is a sequence."""
        s = self.slide("blank", keep=())
        if eyebrow:
            self.eyebrow(s, eyebrow)
        sh.text(s, X0, 1.95, 2.5, 0.9, list(heading), size=26, bold=True, color=self.pal.accent)
        if blurb:
            sh.text(s, X0, 2.95, 2.5, 0.4, blurb, size=9, color=self.pal.muted)
        x, y = 3.25, 1.02
        rh = min(0.6, 3.75 / len(items))
        for i, (t, d, m) in enumerate(items):
            yy = y + i * rh
            sh.text(s, x, yy, 0.5, rh, f"{i + 1:02d}", size=15, bold=True,
                    color=self.pal.accent, anchor="m")
            sh.text(s, x + 0.55, yy, 2.55, rh, t, size=13, bold=True, color=self.pal.text,
                    anchor="m")
            sh.text(s, x + 3.15, yy, 2.25, rh, d, size=9, italic=True, color=self.pal.muted,
                    anchor="m")
            sh.text(s, x + 5.45, yy, 0.86, rh, m, size=9, bold=True, color=self.pal.accent,
                    anchor="m", align="r")
            if i < len(items) - 1:
                sh.shape(s, MSO_SHAPE.RECTANGLE, x, yy + rh - 0.005, 6.31, 0.01,
                         fill=self.pal.card_line)
        return s

    def content(self, title: str, *, accent=None, eyebrow=None, subtitle=None):
        """Titled content slide; compose the body with assets.shapes primitives."""
        s = self.slide("content")
        if eyebrow:
            self.eyebrow(s, eyebrow)
        self.title(s, title, accent)
        if subtitle:
            self.subtitle(s, subtitle)
        return s
